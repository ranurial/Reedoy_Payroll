import os
import csv
import io
import sqlite3
import hashlib
import calendar
import datetime
from functools import wraps

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    send_file,
    abort,
    jsonify,
)


# ============================================================
# OPTIONAL EXPORT LIBRARIES
# ============================================================

try:
    import openpyxl
    from openpyxl import Workbook
except Exception:
    openpyxl = None
    Workbook = None


try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
except Exception:
    SimpleDocTemplate = None
    Paragraph = None
    Spacer = None
    Table = None
    TableStyle = None
    colors = None
    getSampleStyleSheet = None
    A4 = None
    landscape = None


# ============================================================
# APPLICATION
# ============================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "CHANGE-ME-IN-RENDER"
)


DATABASE_URL = os.environ.get(
    "DATABASE_URL",
    ""
).strip()


DB_PATH = os.environ.get(
    "SQLITE_DB",
    "reedoy_payroll.db"
)


DEFAULT_COMPANY_NAME = (
    "REEDOY TEXTILE DYEING PRINTING & FINISHING"
)


MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


# ============================================================
# FIXED DEPARTMENTS
# ============================================================

DEPARTMENTS = [
    "All Departments",
    "General",
    "Printing",
    "Jigar",
    "Wash",
    "Loop",
    "Stanter",
    "Electrical",
    "Accounts",
    "Design",
]


DEPT_BN = {
    "All Departments": "à¦¸à¦•à¦² à¦¬à¦¿à¦­à¦¾à¦—",
    "General": "à¦¸à¦¾à¦§à¦¾à¦°à¦£",
    "Printing": "à¦ªà§à¦°à¦¿à¦¨à§à¦Ÿà¦¿à¦‚",
    "Jigar": "à¦œà¦¿à¦—à¦¾à¦°",
    "Wash": "à¦“à¦¯à¦¼à¦¾à¦¶",
    "Loop": "à¦²à§à¦ª",
    "Stanter": "à¦¸à§à¦Ÿà§à¦¯à¦¾à¦¨à§à¦Ÿà¦¾à¦°",
    "Electrical": "à¦‡à¦²à§‡à¦•à¦Ÿà§à¦°à¦¿à¦•à§à¦¯à¦¾à¦²",
    "Accounts": "à¦…à§à¦¯à¦¾à¦•à¦¾à¦‰à¦¨à§à¦Ÿà¦¸",
    "Design": "à¦¡à¦¿à¦œà¦¾à¦‡à¦¨",
}


def get_worker_departments():
    """
    Return fixed department list.

    This function does not modify the database.
    """
    return DEPARTMENTS.copy()


# ============================================================
# TRANSLATIONS
# ============================================================

LANG = {
        "Quick Links": "দ্রুত লিংক",
    "Workers": "কর্মী",
    "Advances": "অগ্রিম",
    "Department": "বিভাগ",
    "Salary Report": "বেতন রিপোর্ট",
    "Add Advance": "অগ্রিম যোগ করুন",
    "Name": "নাম",
    "Bangla Name": "বাংলা নাম",
    "Designation": "পদবি",
    "Basic Salary": "মূল বেতন",
    "OT Rate/Hour": "ওটি হার/ঘণ্টা",
    "Refreshment": "নাস্তা",
    "Save": "সংরক্ষণ",
    "Cancel": "বাতিল",
    "Search": "অনুসন্ধান",
    "Action": "কার্যক্রম",
    "Edit": "সম্পাদনা",
    "Delete": "মুছুন",
    "Daily worker attendance and overtime": "দৈনিক কর্মী উপস্থিতি ও ওভারটাইম",
    "Worker": "কর্মী",
    "Select Worker": "কর্মী নির্বাচন করুন",
    "Month": "মাস",
    "Year": "বছর",
    "View Attendance": "উপস্থিতি দেখুন",
    "Worker Information": "কর্মীর তথ্য",
    "Worker ID": "কর্মী আইডি",
    "Payroll Month": "বেতন মাস",
    "Daily Attendance": "দৈনিক উপস্থিতি",
    "Day": "দিন",
    "Date": "তারিখ",
    "Status": "অবস্থা",
    "Present": "উপস্থিত",
    "Absent": "অনুপস্থিত",
    "Present Days": "উপস্থিতির দিন",
    "Absent Days": "অনুপস্থিতির দিন",
    "OT Hours": "ওভারটাইম ঘণ্টা",
    "Save Attendance": "উপস্থিতি সংরক্ষণ করুন",
    "Salary Summary": "বেতনের সারসংক্ষেপ",
    "Absent Deduction": "অনুপস্থিতির কর্তন",
    "Earned Basic": "প্রাপ্য মূল বেতন",
    "OT Amount": "ওভারটাইমের টাকা",
    "Nasta": "নাস্তা",
    "Gross Salary": "মোট বেতন",
    "Advance": "অগ্রিম",
    "Net Payable": "নিট প্রদেয়",
    "Select a worker": "একজন কর্মী নির্বাচন করুন",
    "Please select a worker and month to view attendance.": "উপস্থিতি দেখতে একজন কর্মী ও মাস নির্বাচন করুন।",
}


# ============================================================
# DATABASE TYPE
# ============================================================

def is_postgres():
    return bool(
        DATABASE_URL
        and not DATABASE_URL.startswith("sqlite://")
    )


# ============================================================
# POSTGRES CONNECTION POOL
# ============================================================

_PG_POOL = None


def _get_pg_pool():
    global _PG_POOL

    if _PG_POOL is None:

        try:
            import psycopg2
            from psycopg2.pool import ThreadedConnectionPool
        except Exception as e:
            raise RuntimeError(
                "psycopg2 is required for PostgreSQL. "
                "Add psycopg2-binary to requirements.txt."
            ) from e

        minconn = int(
            os.environ.get("PG_POOL_MIN", "1")
        )

        maxconn = int(
            os.environ.get("PG_POOL_MAX", "4")
        )

        _PG_POOL = ThreadedConnectionPool(
            minconn,
            maxconn,
            DATABASE_URL,
            sslmode="require",
            connect_timeout=10,
            keepalives=1,
            keepalives_idle=30,
            keepalives_interval=10,
            keepalives_count=3,
        )

    return _PG_POOL


# ============================================================
# DATABASE CONNECTION
# ============================================================

def db_connect():

    if is_postgres():
        return _get_pg_pool().getconn()

    conn = sqlite3.connect(
        DB_PATH,
        timeout=30,
        check_same_thread=False,
    )

    conn.row_factory = sqlite3.Row

    return conn


def db_release(conn):

    if conn is None:
        return

    if is_postgres():

        try:

            try:
                if getattr(conn, "status", None) != 1:
                    conn.rollback()
            except Exception:
                pass

            _get_pg_pool().putconn(conn)

        except Exception:

            try:
                conn.close()
            except Exception:
                pass

    else:

        try:
            conn.close()
        except Exception:
            pass


# ============================================================
# SQL PLACEHOLDER
# ============================================================

def placeholders(sql):

    if is_postgres():
        return sql.replace("?", "%s")

    return sql


# ============================================================
# GENERIC EXECUTE
# ============================================================

def execute(
    sql,
    params=(),
    fetch=False,
    many=False,
    commit=False,
):

    conn = db_connect()
    cur = conn.cursor()

    try:

        sql2 = placeholders(sql)

        if many:
            cur.executemany(
                sql2,
                params
            )
        else:
            cur.execute(
                sql2,
                params
            )

        rows = (
            cur.fetchall()
            if fetch
            else None
        )

        if commit:
            conn.commit()

        return rows

    except Exception:

        try:
            conn.rollback()
        except Exception:
            pass

        raise

    finally:

        try:
            cur.close()
        except Exception:
            pass

        db_release(conn)


# ============================================================
# ROW CONVERSION
# ============================================================

def row_dict(cur, row):

    if row is None:
        return None

    if hasattr(row, "keys"):
        return dict(row)

    return {
        description[0]: row[index]
        for index, description
        in enumerate(cur.description)
    }


# ============================================================
# FETCH ALL
# ============================================================

def fetch_all(sql, params=()):

    conn = db_connect()
    cur = conn.cursor()

    try:

        cur.execute(
            placeholders(sql),
            params
        )

        rows = cur.fetchall()

        return [
            row_dict(cur, row)
            for row in rows
        ]

    finally:

        try:
            cur.close()
        except Exception:
            pass

        db_release(conn)


# ============================================================
# FETCH ONE
# ============================================================

def fetch_one(sql, params=()):

    rows = fetch_all(
        sql,
        params
    )

    return rows[0] if rows else None


# ============================================================
# SCALAR
# ============================================================

def scalar(
    sql,
    params=(),
    default=0,
):

    row = fetch_one(
        sql,
        params
    )

    if not row:
        return default

    return next(
        iter(row.values())
    )


# ============================================================
# TABLE EXISTS
# ============================================================

def table_exists(name):

    if is_postgres():

        return bool(
            scalar(
                """
                SELECT EXISTS (
                    SELECT 1
                    FROM information_schema.tables
                    WHERE table_schema='public'
                    AND table_name=?
                )
                """,
                (name,),
                False,
            )
        )

    return bool(
        scalar(
            """
            SELECT COUNT(*)
            FROM sqlite_master
            WHERE type='table'
            AND name=?
            """,
            (name,),
            0,
        )
    )


# ============================================================
# TABLE COLUMNS
# ============================================================

def columns(name):

    if not table_exists(name):
        return set()

    if is_postgres():

        rows = fetch_all(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema='public'
            AND table_name=?
            """,
            (name,),
        )

        return {
            row["column_name"]
            for row in rows
        }

    conn = db_connect()
    cur = conn.cursor()

    try:

        cur.execute(
            "PRAGMA table_info(" + name + ")"
        )

        return {
            row[1]
            for row in cur.fetchall()
        }

    finally:

        try:
            cur.close()
        except Exception:
            pass

        db_release(conn)


# ============================================================
# ADD COLUMN
# ============================================================

def add_column(
    name,
    col,
    typ,
):

    if not table_exists(name):
        return

    if col in columns(name):
        return

    execute(
        f"ALTER TABLE {name} ADD COLUMN {col} {typ}",
        commit=True,
    )


# ============================================================
# PASSWORD
# ============================================================

def hash_password(password):

    return hashlib.sha256(
        str(password).encode("utf-8")
    ).hexdigest()


# ============================================================
# CURRENT DATETIME
# ============================================================

def nowstr():

    return datetime.datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )


# ============================================================
# DATABASE INITIALIZATION
# ============================================================

def init_db():
    """
    Create only missing structures.

    IMPORTANT:
    - Never DROP tables.
    - Never TRUNCATE tables.
    - Never delete migrated workers.
    - Never re-import worker records.
    - Existing worker IDs remain untouched.
    """

    conn = db_connect()
    cur = conn.cursor()

    try:

        # ----------------------------------------------------
        # Workers
        # ----------------------------------------------------

        if is_postgres():

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS workers (
                    id SERIAL PRIMARY KEY,
                    name TEXT NOT NULL,
                    basic_salary DOUBLE PRECISION NOT NULL DEFAULT 0,
                    ot_rate DOUBLE PRECISION NOT NULL DEFAULT 0,
                    department TEXT,
                    designation TEXT,
                    refreshment_bill DOUBLE PRECISION NOT NULL DEFAULT 0,
                    bangla_name TEXT
                )
                """
            )

        else:

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS workers (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    basic_salary REAL NOT NULL DEFAULT 0,
                    ot_rate REAL NOT NULL DEFAULT 0,
                    department TEXT,
                    designation TEXT,
                    refreshment_bill REAL NOT NULL DEFAULT 0,
                    bangla_name TEXT
                )
                """
            )

        # ----------------------------------------------------
        # Attendance
        # ----------------------------------------------------

        if is_postgres():

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS attendance (
                    id SERIAL PRIMARY KEY,
                    worker_id INTEGER,
                    month_year TEXT NOT NULL,
                    present_days INTEGER DEFAULT 0,
                    absent_days INTEGER DEFAULT 0,
                    ot_hours DOUBLE PRECISION DEFAULT 0,
                    advance_deduction DOUBLE PRECISION DEFAULT 0
                )
                """
            )

        else:

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS attendance (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    worker_id INTEGER,
                    month_year TEXT NOT NULL,
                    present_days INTEGER DEFAULT 0,
                    absent_days INTEGER DEFAULT 0,
                    ot_hours REAL DEFAULT 0,
                    advance_deduction REAL DEFAULT 0
                )
                """
            )

        # ----------------------------------------------------
        # Daily Attendance
        # ----------------------------------------------------

        if is_postgres():

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS daily_attendance (
                    id SERIAL PRIMARY KEY,
                    worker_id INTEGER,
                    month_year TEXT NOT NULL,
                    day INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'P'
                )
                """
            )

        else:

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS daily_attendance (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    worker_id INTEGER,
                    month_year TEXT NOT NULL,
                    day INTEGER NOT NULL,
                    status TEXT NOT NULL DEFAULT 'P'
                )
                """
            )

        # ----------------------------------------------------
        # Company Settings
        # ----------------------------------------------------

        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS company_settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )
            """
        )

        # ----------------------------------------------------
        # Users
        # ----------------------------------------------------

        if is_postgres():

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id SERIAL PRIMARY KEY,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT,
                    full_name TEXT,
                    role TEXT DEFAULT 'Operator',
                    active INTEGER DEFAULT 1,
                    created_at TEXT,
                    last_login TEXT,
                    password TEXT
                )
                """
            )

        else:

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT,
                    full_name TEXT,
                    role TEXT DEFAULT 'Operator',
                    active INTEGER DEFAULT 1,
                    created_at TEXT,
                    last_login TEXT,
                    password TEXT
                )
                """
            )

        # ----------------------------------------------------
        # Activity Log
        # ----------------------------------------------------

        if is_postgres():

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS activity_log (
                    id SERIAL PRIMARY KEY,
                    username TEXT,
                    action TEXT,
                    log_time TEXT NOT NULL
                )
                """
            )

        else:

            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS activity_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT,
                    action TEXT,
                    log_time TEXT NOT NULL
                )
                """
            )

        conn.commit()

    finally:

        try:
            cur.close()
        except Exception:
            pass

        db_release(conn)

    # ========================================================
    # ADDITIVE MIGRATION ONLY
    # ========================================================

    worker_columns = [
        ("phone", "TEXT"),
        ("address", "TEXT"),
        ("joining_date", "TEXT"),
        ("status", "TEXT"),
        ("bangla_name", "TEXT"),
    ]

    for col, typ in worker_columns:

        try:
            add_column(
                "workers",
                col,
                typ
            )
        except Exception as e:
            app.logger.warning(
                "Worker column migration %s: %r",
                col,
                e
            )

    try:
        add_column(
            "users",
            "password_hash",
            "TEXT"
        )
    except Exception:
        pass

    try:
        add_column(
            "users",
            "password",
            "TEXT"
        )
    except Exception:
        pass

    try:
        add_column(
            "users",
            "last_login",
            "TEXT"
        )
    except Exception:
        pass

    # ========================================================
    # ADMIN USER
    # ========================================================

    try:

        admin = fetch_one(
            """
            SELECT id
            FROM users
            WHERE LOWER(username)=?
            LIMIT 1
            """,
            ("admin",),
        )

        if not admin:

            user_columns = columns(
                "users"
            )

            names = []
            values = []

            data = [
                (
                    "username",
                    "admin"
                ),
                (
                    "password_hash",
                    hash_password("admin123")
                ),
                (
                    "password",
                    hash_password("admin123")
                ),
                (
                    "full_name",
                    "System Administrator"
                ),
                (
                    "role",
                    "Administrator"
                ),
                (
                    "active",
                    1
                ),
                (
                    "created_at",
                    nowstr()
                ),
            ]

            for name, value in data:

                if name in user_columns:

                    names.append(name)
                    values.append(value)

            if names:

                placeholders_list = ",".join(
                    ["?"] * len(values)
                )

                execute(
                    """
                    INSERT INTO users
                    (""" + ",".join(names) + """)
                    VALUES
                    (""" + placeholders_list + """)
                    """,
                    values,
                    commit=True,
                )

    except Exception as e:

        app.logger.warning(
            "Admin initialization error: %r",
            e
        )

    # ========================================================
    # DEFAULT COMPANY SETTINGS
    # ========================================================

    defaults = {
        "company_name": DEFAULT_COMPANY_NAME,
        "company_address": "",
        "company_phone": "",
        "company_email": "",
        "company_logo": "",
    }

    try:

        for key, value in defaults.items():

            if is_postgres():

                execute(
                    """
                    INSERT INTO company_settings
                    (key,value)
                    VALUES (?,?)
                    ON CONFLICT(key)
                    DO NOTHING
                    """,
                    (key, value),
                    commit=True,
                )

            else:

                execute(
                    """
                    INSERT OR IGNORE INTO company_settings
                    (key,value)
                    VALUES (?,?)
                    """,
                    (key, value),
                    commit=True,
                )

    except Exception as e:

        app.logger.warning(
            "Settings initialization error: %r",
            e
        )


# ============================================================
# SETTINGS
# ============================================================

def get_settings():

    try:

        rows = fetch_all(
            """
            SELECT key,value
            FROM company_settings
            """
        )

        result = {}

        for row in rows:
            result[
                row["key"]
            ] = row["value"]

        if "company_name" not in result:
            result["company_name"] = (
                DEFAULT_COMPANY_NAME
            )

        return result

    except Exception:

        return {
            "company_name":
                DEFAULT_COMPANY_NAME
        }


# ============================================================
# ACTIVITY LOG
# ============================================================

def log_activity(
    username,
    action
):

    try:

        execute(
            """
            INSERT INTO activity_log
            (username,action,log_time)
            VALUES (?,?,?)
            """,
            (
                username,
                action,
                nowstr(),
            ),
            commit=True,
        )

    except Exception as e:

        app.logger.warning(
            "Activity log error: %r",
            e
        )


# ============================================================
# CURRENT USER
# ============================================================

def current_user():

    user_id = session.get(
        "user_id"
    )

    if not user_id:
        return None

    try:

        return fetch_one(
            """
            SELECT *
            FROM users
            WHERE id=?
            """,
            (user_id,),
        )

    except Exception:

        return None


# ============================================================
# LOGIN REQUIRED
# ============================================================

def login_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        if not session.get("user_id"):

            return redirect(
                url_for(
                    "login",
                    next=request.path
                )
            )

        return function(
            *args,
            **kwargs
        )

    return wrapper


# ============================================================
# ADMIN REQUIRED
# ============================================================

def admin_required(function):

    @wraps(function)
    def wrapper(*args, **kwargs):

        user = current_user()

        if not user:

            return redirect(
                url_for("login")
            )

        role = str(
            user.get("role", "")
        ).lower()

        if role not in (
            "administrator",
            "admin",
        ):

            flash(
                "Administrator access required.",
                "danger"
            )

            return redirect(
                url_for("dashboard")
            )

        return function(
            *args,
            **kwargs
        )

    return wrapper


# ============================================================
# MONTH
# ============================================================

def month_name_year(
    month=None,
    year=None
):

    month = (
        month
        or request.values.get("month")
        or MONTHS[
            datetime.date.today().month - 1
        ]
    )

    year = (
        year
        or request.values.get("year")
        or str(
            datetime.date.today().year
        )
    )

    if str(month) not in MONTHS:

        month = MONTHS[
            datetime.date.today().month - 1
        ]

    try:
        year = int(year)
    except Exception:
        year = datetime.date.today().year

    return f"{month} {year}"


# ============================================================
# NUMBER PARSER
# ============================================================

def parse_num(
    value,
    default=0.0
):

    try:

        if value is None:
            return default

        text = str(value).strip()

        if text == "":
            return default

        return float(text)

    except Exception:

        return default


# ============================================================
# MONTH FROM DATE
# ============================================================

def month_name_from_date(value):

    try:

        text = str(value)[:10]

        year, month, day = (
            text.split("-")
        )

        return (
            f"{MONTHS[int(month)-1]} "
            f"{int(year)}"
        )

    except Exception:

        return ""


# ============================================================
# LEGACY ADVANCE TABLE
# ============================================================

def legacy_advance_source():

    if table_exists(
        "worker_advances"
    ):

        cols = columns(
            "worker_advances"
        )

        if {
            "worker_id",
            "amount",
        }.issubset(cols):

            return "worker_advances"

    if table_exists(
        "advance_salary"
    ):

        return "advance_salary"

    if table_exists(
        "advances"
    ):

        return "advances"

    return None


# ============================================================
# ADVANCE ROWS
# ============================================================

def advance_rows(
    month=None,
    worker_id=None
):

    table_name = (
        legacy_advance_source()
    )

    if not table_name:
        return []

    table_columns = columns(
        table_name
    )

    # --------------------------------------------------------
    # Modern worker_advances
    # --------------------------------------------------------

    if table_name == "worker_advances":

        where = []
        params = []

        if month:

            where.append(
                "month_year=?"
            )

            params.append(month)

        if worker_id:

            where.append(
                "worker_id=?"
            )

            params.append(worker_id)

        query = f"""
            SELECT
                wa.id,
                wa.worker_id,
                wa.month_year,
                wa.advance_date,
                wa.amount,
                {("wa.note" if "note" in table_columns else "''")} AS note,
                w.name AS worker_name,
                w.bangla_name,
                w.department
            FROM worker_advances wa
            LEFT JOIN workers w
                ON w.id = wa.worker_id
        """

        if where:

            query += (
                " WHERE "
                + " AND ".join(where)
            )

        query += " ORDER BY id DESC"

        return fetch_all(
            query,
            params
        )

    # --------------------------------------------------------
    # Legacy advance_salary
    # --------------------------------------------------------

    if table_name == "advance_salary":

        if "date" in table_columns:

            date_column = "date"

        elif "advance_date" in table_columns:

            date_column = "advance_date"

        else:

            date_column = None

        query = f"""
            SELECT
                a.id,
                a.worker_id,
                {
                    "a." + date_column
                    if date_column
                    else "NULL"
                } AS advance_date,
                a.amount,
                {("a.note" if "note" in table_columns else "''")} AS note,
                w.name AS worker_name,
                w.bangla_name,
                w.department
            FROM advance_salary a
            LEFT JOIN workers w
                ON w.id = a.worker_id
        """

        rows = fetch_all(query)

        for row in rows:

            row["month_year"] = (
                month_name_from_date(
                    row.get(
                        "advance_date"
                    )
                )
            )

        if month:

            rows = [
                row
                for row in rows
                if row.get(
                    "month_year"
                ) == month
            ]

        if worker_id:

            rows = [
                row
                for row in rows
                if str(
                    row.get("worker_id")
                ) == str(worker_id)
            ]

        return rows

    # --------------------------------------------------------
    # Legacy advances
    # --------------------------------------------------------

    date_column = (
        "date"
        if "date" in table_columns
        else (
            "advance_date"
            if "advance_date" in table_columns
            else None
        )
    )
    note_expr = "a.note" if "note" in table_columns else "''"
    date_expr = f"a.{date_column}" if date_column else "NULL"

    query = f"""
        SELECT
            a.id,
            a.worker_id,
            {date_expr} AS advance_date,
            a.amount,
            {note_expr} AS note,
            w.name AS worker_name,
            w.bangla_name,
            w.department
        FROM {table_name} a
        LEFT JOIN workers w
            ON w.id = a.worker_id
    """

    rows = fetch_all(query)

    for row in rows:

        row["month_year"] = (
            month_name_from_date(
                row.get(
                    "advance_date"
                )
            )
        )

    if month:

        rows = [
            row
            for row in rows
            if row.get(
                "month_year"
            ) == month
        ]

    if worker_id:

        rows = [
            row
            for row in rows
            if str(
                row.get("worker_id")
            ) == str(worker_id)
        ]

    return rows


# ============================================================
# SAVE ADVANCE
# ============================================================

def save_advance_record(
    worker_id,
    month,
    advance_date,
    amount,
    note
):

    table_name = (
        legacy_advance_source()
    )

    # --------------------------------------------------------
    # Existing worker_advances
    # --------------------------------------------------------

    if table_name == "worker_advances":

        table_columns = columns(
            "worker_advances"
        )

        names = []
        values = []

        data = [
            (
                "worker_id",
                worker_id
            ),
            (
                "month_year",
                month
            ),
            (
                "advance_date",
                advance_date
            ),
            (
                "amount",
                amount
            ),
            (
                "note",
                note
            ),
        ]

        for name, value in data:

            if name in table_columns:

                names.append(name)
                values.append(value)

        execute(
            """
            INSERT INTO worker_advances
            (""" + ",".join(names) + """)
            VALUES
            (""" + ",".join(
                ["?"] * len(values)
            ) + """)
            """,
            values,
            commit=True,
        )

        return

    # --------------------------------------------------------
    # Existing advance_salary
    # --------------------------------------------------------

    if table_name == "advance_salary":

        table_columns = columns(
            "advance_salary"
        )

        date_column = (
            "date"
            if "date" in table_columns
            else (
                "advance_date"
                if "advance_date"
                in table_columns
                else None
            )
        )

        if not date_column:
            raise RuntimeError(
                "advance_salary has no date column."
            )

        names = [
            "worker_id",
            date_column,
            "amount",
        ]

        values = [
            worker_id,
            advance_date,
            amount,
        ]

        if "note" in table_columns:

            names.append("note")
            values.append(note)

        execute(
            """
            INSERT INTO advance_salary
            (""" + ",".join(names) + """)
            VALUES
            (""" + ",".join(
                ["?"] * len(values)
            ) + """)
            """,
            values,
            commit=True,
        )

        return

    # --------------------------------------------------------
    # Existing advances
    # --------------------------------------------------------

    if table_name == "advances":

        table_columns = columns(
            "advances"
        )

        if (
            "worker_id" in table_columns
            and "date" in table_columns
            and "amount" in table_columns
        ):

            names = [
                "worker_id",
                "date",
                "amount",
            ]

            values = [
                worker_id,
                advance_date,
                amount,
            ]

            if "note" in table_columns:

                names.append("note")
                values.append(note)

            execute(
                """
                INSERT INTO advances
                (""" + ",".join(names) + """)
                VALUES
                (""" + ",".join(
                    ["?"] * len(values)
                ) + """)
                """,
                values,
                commit=True,
            )

            return

    # --------------------------------------------------------
    # No advance table exists.
    # Create only new table.
    # --------------------------------------------------------

    if is_postgres():

        execute(
            """
            CREATE TABLE IF NOT EXISTS worker_advances (
                id SERIAL PRIMARY KEY,
                worker_id INTEGER,
                month_year TEXT NOT NULL,
                advance_date TEXT NOT NULL,
                amount DOUBLE PRECISION NOT NULL,
                note TEXT
            )
            """,
            commit=True,
        )

    else:

        execute(
            """
            CREATE TABLE IF NOT EXISTS worker_advances (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                worker_id INTEGER,
                month_year TEXT NOT NULL,
                advance_date TEXT NOT NULL,
                amount REAL NOT NULL,
                note TEXT
            )
            """,
            commit=True,
        )

    execute(
        """
        INSERT INTO worker_advances
        (
            worker_id,
            month_year,
            advance_date,
            amount,
            note
        )
        VALUES
        (?,?,?,?,?)
        """,
        (
            worker_id,
            month,
            advance_date,
            amount,
            note,
        ),
        commit=True,
    )


# ============================================================
# EDIT ADVANCE
# ============================================================

@app.route(
    "/advance/edit/<int:aid>",
    methods=["GET", "POST"]
)
@app.route(
    "/advances/edit/<int:aid>",
    methods=["GET", "POST"]
)
@login_required
def edit_advance(aid):

    rows = advance_rows()

    record = None

    for row in rows:
        if str(row.get("id")) == str(aid):
            record = row
            break

    if not record:
        flash(
            "Advance record not found.",
            "danger"
        )
        return redirect(
            url_for("advance")
        )

    if request.method == "GET":

        return render_template(
            "advance_edit.html",
            record=record,
            workers=fetch_all(
                """
                SELECT
                    id,
                    name,
                    bangla_name,
                    department
                FROM workers
                ORDER BY id
                """
            ),
        )

    form = request.form

    try:

        worker_id = int(
            form.get("worker_id")
        )

    except Exception:

        flash(
            "Invalid worker.",
            "danger"
        )

        return redirect(
            url_for(
                "edit_advance",
                aid=aid
            )
        )

    month = (
        form.get("month_year")
        or month_name_year(
            form.get("month"),
            form.get("year")
        )
        or record.get("month_year")
    )

    advance_date = (
        form.get("advance_date")
        or record.get("advance_date")
        or datetime.date.today().isoformat()
    )

    amount = parse_num(
        form.get("amount"),
        0
    )

    note = (
        form.get("note", "")
        .strip()
    )

    if amount <= 0:

        flash(
            "Amount must be greater than zero.",
            "danger"
        )

        return redirect(
            url_for(
                "edit_advance",
                aid=aid
            )
        )

    try:

        update_advance_record(
            aid,
            worker_id,
            month,
            advance_date,
            amount,
            note,
        )

        flash(
            "Advance updated successfully.",
            "success"
        )

    except Exception as e:

        app.logger.exception(
            "Advance update error"
        )

        flash(
            f"Could not update advance: {e}",
            "danger"
        )

        return redirect(
            url_for(
                "edit_advance",
                aid=aid
            )
        )

    return redirect(
        url_for(
            "advance",
            month=month
        )
    )


# ============================================================
# DELETE ADVANCE
# ============================================================

def delete_advance_record(
    advance_id
):

    table_name = (
        legacy_advance_source()
    )

    if not table_name:
        return

    execute(
        f"""
        DELETE FROM {table_name}
        WHERE id=?
        """,
        (advance_id,),
        commit=True,
    )


# ============================================================
# UPDATE ADVANCE
# ============================================================

def update_advance_record(
    advance_id,
    worker_id,
    month,
    advance_date,
    amount,
    note
):
    table_name = legacy_advance_source()

    if not table_name:
        raise RuntimeError(
            "No advance table exists."
        )

    table_columns = columns(
        table_name
    )

    if "worker_id" not in table_columns:
        raise RuntimeError(
            f"{table_name} has no worker_id column."
        )

    if "amount" not in table_columns:
        raise RuntimeError(
            f"{table_name} has no amount column."
        )

    # --------------------------------------------------------
    # Modern worker_advances
    # --------------------------------------------------------

    if table_name == "worker_advances":

        fields = []
        values = []

        if "worker_id" in table_columns:
            fields.append("worker_id=?")
            values.append(worker_id)

        if "month_year" in table_columns:
            fields.append("month_year=?")
            values.append(month)

        if "advance_date" in table_columns:
            fields.append("advance_date=?")
            values.append(advance_date)

        if "amount" in table_columns:
            fields.append("amount=?")
            values.append(amount)

        if "note" in table_columns:
            fields.append("note=?")
            values.append(note)

        values.append(advance_id)

        execute(
            """
            UPDATE worker_advances
            SET
            """ + ",".join(fields) + """
            WHERE id=?
            """,
            values,
            commit=True,
        )

        return

    # --------------------------------------------------------
    # Legacy advance_salary
    # --------------------------------------------------------

    if table_name == "advance_salary":

        date_column = (
            "date"
            if "date" in table_columns
            else (
                "advance_date"
                if "advance_date" in table_columns
                else None
            )
        )

        if not date_column:
            raise RuntimeError(
                "advance_salary has no date column."
            )

        fields = [
            "worker_id=?",
            f"{date_column}=?",
            "amount=?",
        ]

        values = [
            worker_id,
            advance_date,
            amount,
        ]

        if "note" in table_columns:
            fields.append("note=?")
            values.append(note)

        values.append(advance_id)

        execute(
            """
            UPDATE advance_salary
            SET
            """ + ",".join(fields) + """
            WHERE id=?
            """,
            values,
            commit=True,
        )

        return

    # --------------------------------------------------------
    # Legacy advances
    # --------------------------------------------------------

    if table_name == "advances":

        date_column = (
            "date"
            if "date" in table_columns
            else (
                "advance_date"
                if "advance_date" in table_columns
                else None
            )
        )

        if not date_column:
            raise RuntimeError(
                "advances has no date column."
            )

        fields = [
            "worker_id=?",
            f"{date_column}=?",
            "amount=?",
        ]

        values = [
            worker_id,
            advance_date,
            amount,
        ]

        if "note" in table_columns:
            fields.append("note=?")
            values.append(note)

        values.append(advance_id)

        execute(
            """
            UPDATE advances
            SET
            """ + ",".join(fields) + """
            WHERE id=?
            """,
            values,
            commit=True,
        )

        return


# ============================================================
# DAILY ATTENDANCE MAP
# ============================================================

def daily_map(
    worker_id,
    month
):

    rows = fetch_all(
        """
        SELECT day,status
        FROM daily_attendance
        WHERE worker_id=?
        AND month_year=?
        """,
        (
            worker_id,
            month,
        ),
    )

    result = {}

    for row in rows:

        try:

            result[
                int(row["day"])
            ] = str(
                row["status"] or "P"
            )

        except Exception:
            pass

    return result


# ============================================================
# BULK SALARY CALCULATION
# ============================================================

def calculate_salary_bulk(
    workers,
    month
):

    worker_ids = [
        worker["id"]
        for worker in workers
    ]

    if not worker_ids:
        return {}

    # --------------------------------------------------------
    # Attendance
    # --------------------------------------------------------

    attendance_rows = fetch_all(
        """
        SELECT *
        FROM attendance
        WHERE month_year=?
        ORDER BY id DESC
        """,
        (month,),
    )

    attendance_by_worker = {}

    for row in attendance_rows:

        worker_id = row.get(
            "worker_id"
        )

        if worker_id not in attendance_by_worker:

            attendance_by_worker[
                worker_id
            ] = row

    # --------------------------------------------------------
    # Daily attendance
    # --------------------------------------------------------

    daily_rows = fetch_all(
        """
        SELECT
            worker_id,
            day,
            status
        FROM daily_attendance
        WHERE month_year=?
        """,
        (month,),
    )

    daily_by_worker = {}

    for row in daily_rows:

        try:

            worker_id = row.get(
                "worker_id"
            )

            day = int(
                row.get("day")
            )

            status = str(
                row.get("status")
                or "P"
            )

            daily_by_worker.setdefault(
                worker_id,
                {}
            )[day] = status

        except Exception:
            pass

    # --------------------------------------------------------
    # Advances
    # --------------------------------------------------------

    advances_by_worker = {}

    table_name = (
        legacy_advance_source()
    )

    if table_name:

        table_columns = columns(
            table_name
        )

        try:

            if (
                table_name
                == "worker_advances"
                and {
                    "worker_id",
                    "amount",
                    "month_year",
                }.issubset(
                    table_columns
                )
            ):

                advance_data = fetch_all(
                    """
                    SELECT
                        worker_id,
                        amount
                    FROM worker_advances
                    WHERE month_year=?
                    """,
                    (month,),
                )

                for row in advance_data:

                    worker_id = row.get(
                        "worker_id"
                    )

                    advances_by_worker[
                        worker_id
                    ] = (
                        advances_by_worker.get(
                            worker_id,
                            0
                        )
                        + parse_num(
                            row.get("amount")
                        )
                    )

            elif (
                table_name
                in (
                    "advance_salary",
                    "advances",
                )
                and "worker_id"
                in table_columns
                and "amount"
                in table_columns
            ):

                if "date" in table_columns:

                    date_column = "date"

                elif (
                    "advance_date"
                    in table_columns
                ):

                    date_column = (
                        "advance_date"
                    )

                else:

                    date_column = None

                if date_column:

                    advance_data = fetch_all(
                        f"""
                        SELECT
                            worker_id,
                            {date_column}
                                AS advance_date,
                            amount
                        FROM {table_name}
                        """
                    )

                    for row in advance_data:

                        if (
                            month_name_from_date(
                                row.get(
                                    "advance_date"
                                )
                            )
                            == month
                        ):

                            worker_id = row.get(
                                "worker_id"
                            )

                            advances_by_worker[
                                worker_id
                            ] = (
                                advances_by_worker.get(
                                    worker_id,
                                    0
                                )
                                + parse_num(
                                    row.get(
                                        "amount"
                                    )
                                )
                            )

        except Exception as e:

            app.logger.warning(
                "Bulk advance load error: %r",
                e
            )

    # --------------------------------------------------------
    # Calculate
    # --------------------------------------------------------

    result = {}

    try:

        month_name, year_text = (
            month.split()
        )

        year = int(year_text)

        month_number = (
            MONTHS.index(
                month_name
            ) + 1
        )

        days_in_month = (
            calendar.monthrange(
                year,
                month_number
            )[1]
        )

    except Exception:

        year = datetime.date.today().year
        month_number = (
            datetime.date.today().month
        )

        days_in_month = 30

    for worker in workers:

        worker_id = worker["id"]

        attendance = (
            attendance_by_worker.get(
                worker_id,
                {}
            )
        )

        present = int(
            attendance.get(
                "present_days"
            )
            or 0
        )

        absent = int(
            attendance.get(
                "absent_days"
            )
            or 0
        )

        overtime = parse_num(
            attendance.get(
                "ot_hours"
            ),
            0
        )

        basic_salary = parse_num(
            worker.get(
                "basic_salary"
            ),
            0
        )

        ot_rate = parse_num(
            worker.get(
                "ot_rate"
            ),
            0
        )

        nasta_rate = parse_num(
            worker.get(
                "refreshment_bill"
            ),
            0
        )

        # ----------------------------------------------------
        # Absent deduction
        # ----------------------------------------------------

        absent_cut = (
            basic_salary
            / days_in_month
            * absent
            if days_in_month
            else 0
        )

        earned_basic = max(
            0,
            basic_salary - absent_cut
        )

        # ----------------------------------------------------
        # OT
        # ----------------------------------------------------

        ot_amount = (
            overtime * ot_rate
        )

        # ----------------------------------------------------
        # Nasta
        # ----------------------------------------------------

        daily = daily_by_worker.get(
            worker_id,
            {}
        )

        if daily:

            billable_days = 0

            for day in range(
                1,
                days_in_month + 1
            ):

                status = daily.get(
                    day,
                    "P"
                )

                try:

                    current_date = (
                        datetime.date(
                            year,
                            month_number,
                            day,
                        )
                    )

                    is_friday = (
                        current_date.weekday()
                        == 4
                    )

                except Exception:

                    is_friday = False

                if (
                    status == "P"
                    and not is_friday
                ):

                    billable_days += 1

            nasta = (
                billable_days
                * nasta_rate
            )

        else:

            non_friday_days = sum(
                1
                for day in range(
                    1,
                    days_in_month + 1
                )
                if datetime.date(
                    year,
                    month_number,
                    day
                ).weekday() != 4
            )

            if days_in_month:

                nasta = round(
                    present
                    * (
                        non_friday_days
                        / days_in_month
                    )
                    * nasta_rate,
                    2,
                )

            else:

                nasta = 0

        # ----------------------------------------------------
        # Advance
        # ----------------------------------------------------

        advance = (
            advances_by_worker.get(
                worker_id,
                0
            )
        )

        # ----------------------------------------------------
        # Gross / Net
        # ----------------------------------------------------

        gross = (
            earned_basic
            + ot_amount
            + nasta
        )

        net = (
            gross
            - advance
        )

        result[worker_id] = {
            "present": present,
            "absent": absent,
            "ot": overtime,
            "absent_cut": absent_cut,
            "earned_basic": earned_basic,
            "ot_amt": ot_amount,
            "nasta": nasta,
            "gross": gross,
            "advance": advance,
            "net": net,
        }

    return result


# ============================================================
# SINGLE SALARY
# ============================================================

def calculate_salary(
    worker,
    month,
    attendance=None
):

    if not worker:
        return {}

    if attendance is None:

        attendance = fetch_one(
            """
            SELECT *
            FROM attendance
            WHERE worker_id=?
            AND month_year=?
            ORDER BY id DESC
            LIMIT 1
            """,
            (
                worker["id"],
                month,
            ),
        ) or {}

    present = int(
        attendance.get(
            "present_days"
        )
        or 0
    )

    absent = int(
        attendance.get(
            "absent_days"
        )
        or 0
    )

    overtime = parse_num(
        attendance.get(
            "ot_hours"
        ),
        0
    )

    try:

        month_name, year_text = (
            month.split()
        )

        year = int(year_text)

        month_number = (
            MONTHS.index(
                month_name
            ) + 1
        )

        days_in_month = (
            calendar.monthrange(
                year,
                month_number
            )[1]
        )

    except Exception:

        year = datetime.date.today().year
        month_number = (
            datetime.date.today().month
        )
        days_in_month = 30

    basic_salary = parse_num(
        worker.get(
            "basic_salary"
        ),
        0
    )

    ot_rate = parse_num(
        worker.get(
            "ot_rate"
        ),
        0
    )

    nasta_rate = parse_num(
        worker.get(
            "refreshment_bill"
        ),
        0
    )

    absent_cut = (
        basic_salary
        / days_in_month
        * absent
        if days_in_month
        else 0
    )

    earned_basic = max(
        0,
        basic_salary - absent_cut
    )

    ot_amount = (
        overtime * ot_rate
    )

    daily = daily_map(
        worker["id"],
        month
    )

    if daily:

        billable_days = 0

        for day in range(
            1,
            days_in_month + 1
        ):

            status = daily.get(
                day,
                "P"
            )

            current_date = datetime.date(
                year,
                month_number,
                day
            )

            if (
                status == "P"
                and current_date.weekday()
                != 4
            ):

                billable_days += 1

        nasta = (
            billable_days
            * nasta_rate
        )

    else:

        non_friday_days = sum(
            1
            for day in range(
                1,
                days_in_month + 1
            )
            if datetime.date(
                year,
                month_number,
                day
            ).weekday() != 4
        )

        nasta = (
            round(
                present
                * (
                    non_friday_days
                    / days_in_month
                )
                * nasta_rate,
                2,
            )
            if days_in_month
            else 0
        )

    advances = sum(
        parse_num(
            row.get("amount")
        )
        for row in advance_rows(
            month,
            worker["id"]
        )
    )

    gross = (
        earned_basic
        + ot_amount
        + nasta
    )

    return {
        "present": present,
        "absent": absent,
        "ot": overtime,
        "absent_cut": absent_cut,
        "earned_basic": earned_basic,
        "ot_amt": ot_amount,
        "nasta": nasta,
        "gross": gross,
        "advance": advances,
        "net": gross - advances,
    }


# ============================================================
# TEMPLATE GLOBALS
# ============================================================

@app.context_processor
def inject_globals():

    settings = get_settings()

    language = session.get(
        "language",
        "en"
    )

    def translate(text):

        if language == "bn":

            return LANG.get(
                text,
                text
            )

        return text

    def department_translate(value):

        value = str(
            value or ""
        )

        if language == "bn":

            return DEPT_BN.get(
                value,
                value
            )

        return value

    return {
        "current_user": current_user(),
        "settings": settings,
        "language": language,
        "tr": translate,
        "dept": department_translate,
        "display_dept": department_translate,
        "months": MONTHS,
        "years": list(
            range(
                2024,
                2032
            )
        ),
        "departments": DEPARTMENTS,
    }


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    try:

        workers_count = scalar(
            "SELECT COUNT(*) FROM workers",
            default=0
        )

        attendance_count = scalar(
            "SELECT COUNT(*) FROM attendance",
            default=0
        )

        daily_count = scalar(
            "SELECT COUNT(*) FROM daily_attendance",
            default=0
        )

        return jsonify({
            "status": "ok",
            "database": (
                "postgresql"
                if is_postgres()
                else "sqlite"
            ),
            "workers": workers_count,
            "attendance": attendance_count,
            "daily_attendance": daily_count,
        })

    except Exception as e:

        return jsonify({
            "status": "error",
            "error": repr(e),
        }), 500


# ============================================================
# DATABASE CHECK
# ============================================================

@app.route("/db-check")
@login_required
def db_check():

    information = {}

    tables = [
        "workers",
        "attendance",
        "daily_attendance",
        "company_settings",
        "users",
        "activity_log",
        "worker_advances",
        "advance_salary",
        "advances",
    ]

    for table_name in tables:

        try:

            exists = table_exists(
                table_name
            )

            information[
                table_name
            ] = {
                "exists": exists,
                "columns": sorted(
                    columns(table_name)
                ) if exists else [],
            }

        except Exception as e:

            information[
                table_name
            ] = {
                "error": repr(e)
            }

    return render_template(
        "db_check.html",
        info=information
    )


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        user = fetch_one(
            """
            SELECT *
            FROM users
            WHERE username=?
            LIMIT 1
            """,
            (username,),
        )

        valid = False

        if user:

            active = user.get(
                "active"
            )

            if active is None:
                active = 1

            if int(active):

                stored_password = (
                    user.get(
                        "password_hash"
                    )
                    or user.get(
                        "password"
                    )
                    or ""
                )

                valid = (
                    stored_password
                    == hash_password(password)
                    or stored_password
                    == password
                )

        if valid:

            session["user_id"] = (
                user["id"]
            )

            session["language"] = (
                session.get(
                    "language",
                    "en"
                )
            )

            try:

                if "last_login" in columns(
                    "users"
                ):

                    execute(
                        """
                        UPDATE users
                        SET last_login=?
                        WHERE id=?
                        """,
                        (
                            nowstr(),
                            user["id"],
                        ),
                        commit=True,
                    )

            except Exception:
                pass

            log_activity(
                username,
                "Logged in"
            )

            next_url = request.args.get(
                "next"
            )

            if next_url:
                return redirect(next_url)

            return redirect(
                url_for("dashboard")
            )

        flash(
            "Invalid User ID or Password.",
            "danger"
        )

    return render_template(
        "login.html"
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    user = current_user()

    if user:

        log_activity(
            user.get("username"),
            "Logged out"
        )

    session.clear()

    return redirect(
        url_for("login")
    )


# ============================================================
# LANGUAGE
# ============================================================

@app.route("/language/<lang>")
def language(lang):

    if lang == "bn":

        session["language"] = "bn"

    else:

        session["language"] = "en"

    return redirect(
        request.referrer
        or url_for("dashboard")
    )


# ============================================================
# INDEX
# ============================================================

@app.route("/")
@login_required
def index():

    return redirect(
        url_for("dashboard")
    )


# ============================================================
# DASHBOARD
# ============================================================

@app.route("/dashboard")
@login_required
def dashboard():

    month = month_name_year()

    workers = fetch_all(
        """
        SELECT *
        FROM workers
        ORDER BY id
        """
    )

    salaries = calculate_salary_bulk(
        workers,
        month
    )

    gross = sum(
        salary["gross"]
        for salary in salaries.values()
    )

    advance = sum(
        salary["advance"]
        for salary in salaries.values()
    )

    today = datetime.date.today()

    today_month = (
        f"{MONTHS[today.month - 1]} "
        f"{today.year}"
    )

    today_rows = fetch_all(
        """
        SELECT
            status,
            COUNT(*) AS c
        FROM daily_attendance
        WHERE month_year=?
        AND day=?
        GROUP BY status
        """,
        (
            today_month,
            today.day,
        ),
    )

    today_map = {
        row["status"]: row["c"]
        for row in today_rows
    }

    department_rows = fetch_all(
        """
        SELECT
            department,
            COUNT(*) AS worker_count,
            COALESCE(
                SUM(basic_salary),
                0
            ) AS total_basic
        FROM workers
        GROUP BY department
        ORDER BY department
        """
    )

    return render_template(
        "dashboard.html",
        workers=workers,
        month=month,
        month_name=month.split()[0],
        year=month.split()[1],
        gross=gross,
        advance=advance,
        today_present=today_map.get(
            "P",
            0
        ),
        today_absent=today_map.get(
            "A",
            0
        ),
        dept_rows=department_rows,
        total_workers=len(workers),
    )


# ============================================================
# WORKER LIST
# ============================================================

@app.route("/workers")
@login_required
def workers():

    query_text = request.args.get(
        "q",
        ""
    ).strip()

    params = []

    sql = """
        SELECT *
        FROM workers
    """

    if query_text:

        like_value = (
            f"%{query_text.lower()}%"
        )

        sql += """
            WHERE
                CAST(id AS TEXT)
                    LIKE ?
                OR LOWER(
                    COALESCE(name, '')
                ) LIKE ?
                OR LOWER(
                    COALESCE(bangla_name, '')
                ) LIKE ?
                OR LOWER(
                    COALESCE(department, '')
                ) LIKE ?
                OR LOWER(
                    COALESCE(designation, '')
                ) LIKE ?
        """

        params = [
            like_value,
            like_value,
            like_value,
            like_value,
            like_value,
        ]

    sql += """
        ORDER BY id DESC
    """

    worker_rows = fetch_all(
        sql,
        params
    )

    return render_template(
        "worker.html",
        workers=worker_rows,
        q=query_text,
    )


# ============================================================
# ADD WORKER
# ============================================================

@app.route(
    "/workers/add",
    methods=["GET", "POST"]
)
@login_required
def add_worker():

    if request.method == "POST":

        form = request.form

        name = form.get(
            "name",
            ""
        ).strip()

        bangla_name = form.get(
            "bangla_name",
            ""
        ).strip()

        department = form.get(
            "department",
            ""
        ).strip()

        designation = form.get(
            "designation",
            ""
        ).strip()

        basic_salary = parse_num(
            form.get("basic_salary"),
            0
        )

        ot_rate = parse_num(
            form.get("ot_rate"),
            0
        )

        refreshment_bill = parse_num(
            form.get(
                "refreshment_bill"
            ),
            0
        )

        if not name:

            flash(
                "Worker name is required.",
                "danger"
            )

            return render_template(
                "worker_form.html",
                worker=form,
                departments=(
                    get_worker_departments()
                ),
                form_title="Add Worker",
                submit_text="Save",
            )

        if (
            department
            not in DEPARTMENTS
            or department
            == "All Departments"
        ):

            department = "General"

        try:

            execute(
                """
                INSERT INTO workers
                (
                    name,
                    basic_salary,
                    ot_rate,
                    department,
                    designation,
                    refreshment_bill,
                    bangla_name
                )
                VALUES
                (?,?,?,?,?,?,?)
                """,
                (
                    name,
                    basic_salary,
                    ot_rate,
                    department,
                    designation,
                    refreshment_bill,
                    bangla_name,
                ),
                commit=True,
            )

            flash(
                "Worker saved successfully.",
                "success"
            )

            return redirect(
                url_for("workers")
            )

        except Exception as e:

            app.logger.exception(
                "Worker insert error"
            )

            flash(
                f"Could not save worker: {e}",
                "danger"
            )

            return render_template(
                "worker_form.html",
                worker=form,
                departments=(
                    get_worker_departments()
                ),
                form_title="Add Worker",
                submit_text="Save",
            )

    return render_template(
        "worker_form.html",
        worker=None,
        departments=(
            get_worker_departments()
        ),
        form_title="Add Worker",
        submit_text="Save",
    )


# ============================================================
# EDIT WORKER
# ============================================================

@app.route(
    "/workers/edit/<int:worker_id>",
    methods=["GET", "POST"]
)
@login_required
def edit_worker(worker_id):

    worker = fetch_one(
        """
        SELECT *
        FROM workers
        WHERE id=?
        """,
        (worker_id,),
    )

    if not worker:

        abort(404)

    if request.method == "POST":

        form = request.form

        name = form.get(
            "name",
            ""
        ).strip()

        bangla_name = form.get(
            "bangla_name",
            ""
        ).strip()

        department = form.get(
            "department",
            ""
        ).strip()

        designation = form.get(
            "designation",
            ""
        ).strip()

        basic_salary = parse_num(
            form.get("basic_salary"),
            0
        )

        ot_rate = parse_num(
            form.get("ot_rate"),
            0
        )

        refreshment_bill = parse_num(
            form.get(
                "refreshment_bill"
            ),
            0
        )

        if not name:

            flash(
                "Worker name is required.",
                "danger"
            )

            return render_template(
                "worker_form.html",
                worker=form,
                departments=(
                    get_worker_departments()
                ),
                form_title="Edit Worker",
                submit_text="Update",
            )

        if (
            department
            not in DEPARTMENTS
            or department
            == "All Departments"
        ):

            department = "General"

        try:

            execute(
                """
                UPDATE workers
                SET
                    name=?,
                    basic_salary=?,
                    ot_rate=?,
                    department=?,
                    designation=?,
                    refreshment_bill=?,
                    bangla_name=?
                WHERE id=?
                """,
                (
                    name,
                    basic_salary,
                    ot_rate,
                    department,
                    designation,
                    refreshment_bill,
                    bangla_name,
                    worker_id,
                ),
                commit=True,
            )

            flash(
                "Worker updated successfully.",
                "success"
            )

            return redirect(
                url_for("workers")
            )

        except Exception as e:

            app.logger.exception(
                "Worker update error"
            )

            flash(
                f"Could not update worker: {e}",
                "danger"
            )

            return render_template(
                "worker_form.html",
                worker=form,
                departments=(
                    get_worker_departments()
                ),
                form_title="Edit Worker",
                submit_text="Update",
            )

    return render_template(
        "worker_form.html",
        worker=worker,
        departments=(
            get_worker_departments()
        ),
        form_title="Edit Worker",
        submit_text="Update",
    )


# ============================================================
# DELETE WORKER
# ============================================================

@app.route(
    "/workers/delete/<int:worker_id>",
    methods=["POST", "GET"]
)
@login_required
def delete_worker(worker_id):

    worker = fetch_one(
        """
        SELECT id,name
        FROM workers
        WHERE id=?
        """,
        (worker_id,),
    )

    if not worker:

        flash(
            "Worker not found.",
            "danger"
        )

        return redirect(
            url_for("workers")
        )

    try:

        # Delete related records ONLY
        # after explicit worker deletion.

        for table_name in [
            "attendance",
            "daily_attendance",
            "worker_advances",
            "advance_salary",
            "advances",
        ]:

            if table_exists(
                table_name
            ):

                try:

                    if "worker_id" in columns(
                        table_name
                    ):

                        execute(
                            f"""
                            DELETE FROM {table_name}
                            WHERE worker_id=?
                            """,
                            (worker_id,),
                            commit=True,
                        )

                except Exception as e:

                    app.logger.warning(
                        "Related delete %s: %r",
                        table_name,
                        e
                    )

        execute(
            """
            DELETE FROM workers
            WHERE id=?
            """,
            (worker_id,),
            commit=True,
        )

        flash(
            "Worker deleted successfully.",
            "success"
        )

    except Exception as e:

        app.logger.exception(
            "Worker delete error"
        )

        flash(
            f"Could not delete worker: {e}",
            "danger"
        )

    return redirect(
        url_for("workers")
    )


# ============================================================
# ATTENDANCE
# ============================================================

@app.route("/attendance")
@login_required
def attendance():

    requested_month = request.args.get(
        "month"
    )

    requested_year = request.args.get(
        "year"
    )

    if (
        requested_month
        and requested_year
    ):

        month = month_name_year(
            requested_month,
            requested_year
        )

    else:

        month = month_name_year()

    worker_id = request.args.get(
        "worker_id",
        ""
    )

    worker = None
    days = []
    summary = {}

    if worker_id:

        worker = fetch_one(
            """
            SELECT *
            FROM workers
            WHERE id=?
            """,
            (worker_id,),
        )

        if worker:

            daily = daily_map(
                worker["id"],
                month
            )

            month_name, year_text = (
                month.split()
            )

            year = int(
                year_text
            )

            month_number = (
                MONTHS.index(
                    month_name
                ) + 1
            )

            number_of_days = (
                calendar.monthrange(
                    year,
                    month_number
                )[1]
            )

            days = [
                {
                    "day": day,
                    "status": daily.get(
                        day,
                        "P"
                    ),
                    "date": datetime.date(
                        year,
                        month_number,
                        day
                    ),
                }
                for day in range(
                    1,
                    number_of_days + 1
                )
            ]

            summary = calculate_salary(
                worker,
                month
            )

    return render_template(
        "attendance.html",
        workers=fetch_all(
            """
            SELECT
                id,
                name,
                bangla_name,
                department
            FROM workers
            ORDER BY id
            """
        ),
        worker=worker,
        days=days,
        summary=summary,
        month=month,
    )


# ============================================================
# SAVE ATTENDANCE
# ============================================================

@app.route(
    "/attendance/save",
    methods=["POST"]
)
@login_required
def save_attendance():

    form = request.form

    try:

        worker_id = int(
            form.get("worker_id")
        )

    except Exception:

        flash(
            "Invalid worker.",
            "danger"
        )

        return redirect(
            url_for("attendance")
        )

    if form.get("year"):

        month = month_name_year(
            form.get("month"),
            form.get("year")
        )

    else:

        month = (
            form.get("month_year")
            or month_name_year()
        )

    present = int(
        form.get(
            "present_days",
            0
        )
        or 0
    )

    absent = int(
        form.get(
            "absent_days",
            0
        )
        or 0
    )

    overtime = parse_num(
        form.get("ot_hours"),
        0
    )

    daily_statuses = {}

    raw_json = form.get(
        "statuses_json"
    )

    if raw_json:

        try:

            import json

            daily_statuses = (
                json.loads(raw_json)
            )

        except Exception:

            daily_statuses = {}

    # Accept day_1=P etc.

    for key, value in form.items():

        if key.startswith(
            "day_"
        ):

            daily_statuses[
                key[4:]
            ] = value

    # --------------------------------------------------------
    # Save daily attendance
    # --------------------------------------------------------

    for day, status in (
        daily_statuses.items()
    ):

        try:

            day_number = int(day)

            status = (
                "A"
                if str(status)
                .upper()
                .startswith("A")
                else "P"
            )

            existing = fetch_one(
                """
                SELECT id
                FROM daily_attendance
                WHERE worker_id=?
                AND month_year=?
                AND day=?
                ORDER BY id DESC
                LIMIT 1
                """,
                (
                    worker_id,
                    month,
                    day_number,
                ),
            )

            if existing:

                execute(
                    """
                    UPDATE daily_attendance
                    SET status=?
                    WHERE id=?
                    """,
                    (
                        status,
                        existing["id"],
                    ),
                    commit=True,
                )

            else:

                execute(
                    """
                    INSERT INTO daily_attendance
                    (
                        worker_id,
                        month_year,
                        day,
                        status
                    )
                    VALUES
                    (?,?,?,?)
                    """,
                    (
                        worker_id,
                        month,
                        day_number,
                        status,
                    ),
                    commit=True,
                )

        except Exception as e:

            app.logger.warning(
                "Daily attendance error: %r",
                e
            )

    # --------------------------------------------------------
    # Monthly attendance
    # --------------------------------------------------------

    existing = fetch_one(
        """
        SELECT id
        FROM attendance
        WHERE worker_id=?
        AND month_year=?
        ORDER BY id DESC
        LIMIT 1
        """,
        (
            worker_id,
            month,
        ),
    )

    if existing:

        execute(
            """
            UPDATE attendance
            SET
                present_days=?,
                absent_days=?,
                ot_hours=?
            WHERE id=?
            """,
            (
                present,
                absent,
                overtime,
                existing["id"],
            ),
            commit=True,
        )

    else:

        attendance_columns = columns(
            "attendance"
        )

        names = [
            "worker_id",
            "month_year",
            "present_days",
            "absent_days",
            "ot_hours",
        ]

        values = [
            worker_id,
            month,
            present,
            absent,
            overtime,
        ]

        if (
            "advance_deduction"
            in attendance_columns
        ):

            names.append(
                "advance_deduction"
            )

            values.append(0)

        execute(
            """
            INSERT INTO attendance
            (""" + ",".join(names) + """)
            VALUES
            (""" + ",".join(
                ["?"] * len(values)
            ) + """)
            """,
            values,
            commit=True,
        )

    flash(
        "Attendance saved successfully.",
        "success"
    )

    return redirect(
        url_for(
            "attendance",
            worker_id=worker_id,
            month=month.split()[0],
            year=month.split()[1],
        )
    )


# ============================================================
# PAYSLIP
# ============================================================

@app.route("/payslip")
@login_required
def payslip():

    month = month_name_year()

    worker_id = request.args.get(
        "worker_id"
    )

    worker = (
        fetch_one(
            """
            SELECT *
            FROM workers
            WHERE id=?
            """,
            (worker_id,),
        )
        if worker_id
        else None
    )

    summary = (
        calculate_salary(
            worker,
            month
        )
        if worker
        else {}
    )

    return render_template(
        "payslip.html",
        workers=fetch_all(
            """
            SELECT
                id,
                name,
                bangla_name,
                department
            FROM workers
            ORDER BY id
            """
        ),
        worker=worker,
        summary=summary,
        month=month,
    )


# ============================================================
# DEPARTMENT SALARY
# ============================================================

@app.route("/department")
@login_required
def department():

    month = month_name_year()

    department_name = request.args.get(
        "department",
        ""
    )

    if department_name:

        workers_list = fetch_all(
            """
            SELECT *
            FROM workers
            WHERE department=?
            ORDER BY id
            """,
            (department_name,),
        )

    else:

        workers_list = fetch_all(
            """
            SELECT *
            FROM workers
            ORDER BY id
            """
        )

    salary_map = (
        calculate_salary_bulk(
            workers_list,
            month
        )
    )

    rows = []

    for worker in workers_list:

        row = dict(worker)

        row.update(
            salary_map.get(
                worker["id"],
                {}
            )
        )

        rows.append(row)

    department_list = fetch_all(
        """
        SELECT DISTINCT department
        FROM workers
        WHERE department IS NOT NULL
        AND department <> ''
        ORDER BY department
        """
    )

    return render_template(
        "department.html",
        rows=rows,
        month=month,
        department=department_name,
        departments=[
            row["department"]
            for row in department_list
        ],
    )


# ============================================================
# REPORT COMPATIBILITY
# ============================================================

@app.route("/report", methods=["GET"])
@login_required
def report():
    return department()


# ============================================================
# ADVANCE
# ============================================================

@app.route("/advance")
@app.route("/advances")
@login_required
def advance():

    month = month_name_year()

    worker_id = request.args.get(
        "worker_id"
    )

    return render_template(
        "advance.html",
        workers=fetch_all(
            """
            SELECT
                id,
                name,
                bangla_name,
                department
            FROM workers
            ORDER BY id
            """
        ),
        rows=advance_rows(
            month,
            worker_id
        ),
        month=month,
    )


# Compatibility endpoint.

app.add_url_rule(
    "/advances",
    endpoint="advances",
    view_func=advance,
    methods=["GET"]
)


# ============================================================
# SAVE ADVANCE
# ============================================================

@app.route(
    "/advance/save",
    methods=["POST"]
)
@app.route(
    "/advances/save",
    methods=["POST"]
)
@login_required
def save_advance():

    form = request.form

    try:

        worker_id = int(
            form.get("worker_id")
        )

    except Exception:

        flash(
            "Invalid worker.",
            "danger"
        )

        return redirect(
            url_for("advance")
        )

    month = (
        form.get("month_year")
        or month_name_year(
            form.get("month"),
            form.get("year")
        )
    )

    advance_date = (
        form.get("advance_date")
        or datetime.date.today()
        .isoformat()
    )

    amount = parse_num(
        form.get("amount"),
        0
    )

    note = form.get(
        "note",
        ""
    ).strip()

    if amount <= 0:

        flash(
            "Amount must be greater than zero.",
            "danger"
        )

        return redirect(
            url_for(
                "advance",
                month=month
            )
        )

    try:

        save_advance_record(
            worker_id,
            month,
            advance_date,
            amount,
            note,
        )

        flash(
            "Advance saved successfully.",
            "success"
        )

    except Exception as e:

        app.logger.exception(
            "Advance save error"
        )

        flash(
            f"Could not save advance: {e}",
            "danger"
        )

    return redirect(
        url_for(
            "advance",
            month=month
        )
    )




# ============================================================
# DELETE ADVANCE
# ============================================================

@app.route(
    "/advance/delete/<int:aid>",
    methods=["GET", "POST"]
)
@app.route(
    "/advances/delete/<int:aid>",
    methods=["GET", "POST"]
)
@login_required
def delete_advance(aid):

    try:

        delete_advance_record(
            aid
        )

        flash(
            "Advance deleted successfully.",
            "success"
        )

    except Exception as e:

        app.logger.exception(
            "Advance delete error"
        )

        flash(
            f"Could not delete advance: {e}",
            "danger"
        )

    return redirect(
        url_for("advance")
    )


# ============================================================
# SETTINGS
# ============================================================

@app.route(
    "/settings",
    methods=["GET", "POST"]
)
@admin_required
def settings():

    if request.method == "POST":

        for key in [
            "company_name",
            "company_address",
            "company_phone",
            "company_email",
            "company_logo",
        ]:

            value = request.form.get(
                key,
                ""
            )

            if is_postgres():

                execute(
                    """
                    INSERT INTO company_settings
                    (key,value)
                    VALUES (?,?)
                    ON CONFLICT(key)
                    DO UPDATE
                    SET value=EXCLUDED.value
                    """,
                    (
                        key,
                        value,
                    ),
                    commit=True,
                )

            else:

                execute(
                    """
                    INSERT OR REPLACE
                    INTO company_settings
                    (key,value)
                    VALUES (?,?)
                    """,
                    (
                        key,
                        value,
                    ),
                    commit=True,
                )

        flash(
            "Settings saved successfully.",
            "success"
        )

    return render_template(
        "settings.html",
        settings=get_settings()
    )


# ============================================================
# USERS
# ============================================================

@app.route("/users")
@admin_required
def users():

    return render_template(
        "users.html",
        users=fetch_all(
            """
            SELECT
                id,
                username,
                full_name,
                role,
                active,
                created_at,
                last_login
            FROM users
            ORDER BY id
            """
        ),
    )


# ============================================================
# ADD USER
# ============================================================

@app.route(
    "/users/add",
    methods=["GET", "POST"]
)
@admin_required
def add_user():

    if request.method == "POST":

        form = request.form

        username = form.get(
            "username",
            ""
        ).strip()

        password = form.get(
            "password",
            ""
        )

        if not username or not password:

            flash(
                "Username and password are required.",
                "danger"
            )

            return render_template(
                "user_form.html",
                user=form
            )

        password_hash = (
            hash_password(password)
        )

        user_columns = columns(
            "users"
        )

        names = []
        values = []

        data = [
            (
                "username",
                username
            ),
            (
                "password_hash",
                password_hash
            ),
            (
                "password",
                password_hash
            ),
            (
                "full_name",
                form.get(
                    "full_name",
                    ""
                )
            ),
            (
                "role",
                form.get(
                    "role",
                    "Operator"
                )
            ),
            (
                "active",
                1
            ),
            (
                "created_at",
                nowstr()
            ),
        ]

        for name, value in data:

            if name in user_columns:

                names.append(name)
                values.append(value)

        try:

            execute(
                """
                INSERT INTO users
                (""" + ",".join(names) + """)
                VALUES
                (""" + ",".join(
                    ["?"] * len(values)
                ) + """)
                """,
                values,
                commit=True,
            )

            flash(
                "User created successfully.",
                "success"
            )

            return redirect(
                url_for("users")
            )

        except Exception as e:

            app.logger.exception(
                "User creation error"
            )

            flash(
                f"Could not create user: {e}",
                "danger"
            )

    return render_template(
        "user_form.html",
        user=None
    )


# ============================================================
# DELETE USER
# ============================================================

@app.route(
    "/users/delete/<int:user_id>"
)
@admin_required
def delete_user(user_id):

    user = current_user()

    if (
        user
        and user.get("id")
        == user_id
    ):

        flash(
            "You cannot delete the logged-in user.",
            "danger"
        )

    else:

        execute(
            """
            DELETE FROM users
            WHERE id=?
            """,
            (user_id,),
            commit=True,
        )

        flash(
            "User deleted successfully.",
            "success"
        )

    return redirect(
        url_for("users")
    )


# ============================================================
# ACTIVITY
# ============================================================

@app.route("/activity")
@admin_required
def activity():

    return render_template(
        "activity.html",
        logs=fetch_all(
            """
            SELECT *
            FROM activity_log
            ORDER BY id DESC
            LIMIT 500
            """
        ),
    )


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_rows(
    rows,
    filename,
    title="Reedoy Payroll"
):

    if openpyxl is None:

        abort(
            503,
            description=(
                "openpyxl is not installed."
            )
        )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Payroll"

    if rows:

        headers = list(
            rows[0].keys()
        )

        worksheet.append(
            headers
        )

        for row in rows:

            worksheet.append(
                [
                    row.get(header)
                    for header in headers
                ]
            )

    else:

        worksheet.append(
            ["No records"]
        )

    output = io.BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


# ============================================================
# WORKERS EXCEL
# ============================================================

@app.route(
    "/export/workers.xlsx"
)
@login_required
def export_workers():

    return export_rows(
        fetch_all(
            """
            SELECT *
            FROM workers
            ORDER BY id
            """
        ),
        "reedoy_workers.xlsx"
    )


# ============================================================
# ADVANCES EXCEL
# ============================================================

@app.route(
    "/export/advances.xlsx"
)
@login_required
def export_advances():

    return export_rows(
        advance_rows(
            month_name_year()
        ),
        "reedoy_advances.xlsx"
    )


# ============================================================
# DEPARTMENT EXCEL
# ============================================================

@app.route(
    "/export/department.xlsx"
)
@login_required
def export_department():

    month = month_name_year()

    department_name = request.args.get(
        "department",
        ""
    )

    if department_name:

        worker_list = fetch_all(
            """
            SELECT *
            FROM workers
            WHERE department=?
            ORDER BY id
            """,
            (department_name,),
        )

    else:

        worker_list = fetch_all(
            """
            SELECT *
            FROM workers
            ORDER BY id
            """
        )

    salary_map = (
        calculate_salary_bulk(
            worker_list,
            month
        )
    )

    rows = []

    for worker in worker_list:

        salary = salary_map.get(
            worker["id"],
            {}
        )

        rows.append({
            "ID": worker["id"],
            "Worker Name": worker["name"],
            "Department": worker.get(
                "department"
            ),
            "Designation": worker.get(
                "designation"
            ),
            "Present": salary.get(
                "present",
                0
            ),
            "Absent": salary.get(
                "absent",
                0
            ),
            "OT": salary.get(
                "ot",
                0
            ),
            "Absent Deduction": salary.get(
                "absent_cut",
                0
            ),
            "OT Amount": salary.get(
                "ot_amt",
                0
            ),
            "Nasta": salary.get(
                "nasta",
                0
            ),
            "Gross": salary.get(
                "gross",
                0
            ),
            "Advance": salary.get(
                "advance",
                0
            ),
            "Net Payable": salary.get(
                "net",
                0
            ),
        })

    return export_rows(
        rows,
        "reedoy_department_salary.xlsx"
    )


# ============================================================
# PAYSLIP PDF
# ============================================================

@app.route(
    "/export/payslip.pdf"
)
@login_required
def export_payslip_pdf():

    if (
        SimpleDocTemplate is None
        or Paragraph is None
        or Table is None
        or TableStyle is None
    ):

        abort(
            503,
            description=(
                "reportlab is not installed."
            )
        )

    worker_id = request.args.get(
        "worker_id"
    )

    month = month_name_year()

    worker = fetch_one(
        """
        SELECT *
        FROM workers
        WHERE id=?
        """,
        (worker_id,),
    )

    if not worker:
        abort(404)

    salary = calculate_salary(
        worker,
        month
    )

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=35,
        leftMargin=35,
        topMargin=35,
        bottomMargin=35,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph(
            DEFAULT_COMPANY_NAME,
            styles["Title"]
        ),
        Paragraph(
            "Payroll Payslip - " + month,
            styles["Heading2"]
        ),
        Spacer(
            1,
            12
        ),
    ]

    data = [
        [
            "Worker ID",
            worker["id"]
        ],
        [
            "Worker Name",
            worker["name"]
        ],
        [
            "Department",
            worker.get(
                "department"
            ) or ""
        ],
        [
            "Basic Salary",
            f"BDT {salary['earned_basic']:,.2f}"
        ],
        [
            "Absent Deduction",
            f"BDT {salary['absent_cut']:,.2f}"
        ],
        [
            "OT",
            f"BDT {salary['ot_amt']:,.2f}"
        ],
        [
            "Nasta",
            f"BDT {salary['nasta']:,.2f}"
        ],
        [
            "Gross",
            f"BDT {salary['gross']:,.2f}"
        ],
        [
            "Advance",
            f"BDT {salary['advance']:,.2f}"
        ],
        [
            "Net Payable",
            f"BDT {salary['net']:,.2f}"
        ],
    ]

    story.append(
        Table(
            data,
            colWidths=[
                180,
                280
            ],
            style=TableStyle([
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica"
                ),
                (
                    "PADDING",
                    (0, 0),
                    (-1, -1),
                    6
                ),
            ]),
        )
    )

    document.build(
        story
    )

    output.seek(0)

    filename = (
        f"payslip_"
        f"{worker['id']}_"
        f"{month.replace(' ', '_')}.pdf"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


# ============================================================
# WORKERS CSV
# ============================================================

@app.route(
    "/export/workers.csv"
)
@login_required
def export_workers_csv():

    rows = fetch_all(
        """
        SELECT *
        FROM workers
        ORDER BY id
        """
    )

    output = io.StringIO()

    if rows:

        writer = csv.DictWriter(
            output,
            fieldnames=list(
                rows[0].keys()
            )
        )

        writer.writeheader()

        writer.writerows(
            rows
        )

    else:

        output.write(
            "No records\n"
        )

    data = io.BytesIO(
        output.getvalue().encode(
            "utf-8-sig"
        )
    )

    return send_file(
        data,
        as_attachment=True,
        download_name=(
            "reedoy_workers.csv"
        ),
        mimetype="text/csv",
    )


# ============================================================
# 404
# ============================================================

@app.errorhandler(404)
def not_found(error):

    try:

        return (
            render_template(
                "404.html"
            ),
            404,
        )

    except Exception:

        return (
            "404 - Page not found",
            404,
        )


# ============================================================
# 500
# ============================================================

@app.errorhandler(500)
def server_error(error):

    app.logger.exception(
        "Unhandled application error"
    )

    try:

        return (
            render_template(
                "500.html",
                error=error
            ),
            500,
        )

    except Exception:

        return (
            f"500 - Internal Server Error\n"
            f"{error}",
            500,
        )


# ============================================================
# DATABASE INITIALIZATION
# ============================================================

try:

    init_db()

except Exception as e:

    app.logger.exception(
        "Database initialization error: %r",
        e
    )


# ============================================================
# LOCAL DEVELOPMENT
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=False,
    )





