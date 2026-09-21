import os, csv, io, sqlite3, hashlib, calendar, datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, abort, jsonify

# Optional exports
try:
    import openpyxl
    from openpyxl import Workbook
except Exception:
    openpyxl = None
try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
except Exception:
    SimpleDocTemplate = Paragraph = Spacer = Table = TableStyle = None
    A4 = landscape = None

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "CHANGE-ME-IN-RENDER")
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
DB_PATH = os.environ.get("SQLITE_DB", "reedoy_payroll.db")
DEFAULT_COMPANY_NAME = "REEDOY TEXTILE DYEING PRINTING & FINISHING"

MONTHS = ["January","February","March","April","May","June","July","August","September","October","November","December"]
DEPT_BN = {
    "All Departments":"সব বিভাগ", "Cutting":"কাটিং", "Sewing":"সেলাই", "Finishing":"ফিনিশিং",
    "Quality Control":"কোয়ালিটি কন্ট্রোল", "Knitting":"নিটিং", "Dyeing":"ডাইং", "Printing":"প্রিন্টিং",
    "Washing":"ওয়াশিং", "Iron & Packing":"আয়রন ও প্যাকিং", "Maintenance":"মেইনটেন্যান্স",
    "Store & Logistics":"স্টোর ও লজিস্টিকস", "Human Resources & Admin":"মানবসম্পদ ও প্রশাসন",
    "Accounts & Commercial":"অ্যাকাউন্টস ও কমার্শিয়াল", "Utility":"ইউটিলিটি"
}
LANG = {
    "Dashboard":"ড্যাশবোর্ড", "Workers Management":"কর্মী ব্যবস্থাপনা", "Attendance & Calendar":"উপস্থিতি ও ক্যালেন্ডার",
    "Single Payslip":"একক পে-স্লিপ", "Advance Salary":"অগ্রিম বেতন", "Department Salary Sheet":"বিভাগভিত্তিক বেতন শীট",
    "Settings":"সেটিংস", "Worker Name":"কর্মীর নাম", "Department":"বিভাগ", "Designation":"পদবি",
    "Basic Salary":"মূল বেতন", "OT Rate":"OT হার", "Nasta Rate":"নাস্তা হার", "Present":"উপস্থিত", "Absent":"অনুপস্থিত",
    "Absent Deduction":"অনুপস্থিতির কর্তন", "OT Amt":"OT টাকা", "Nasta":"নাস্তা", "Gross Salary":"মোট বেতন",
    "Advance":"অগ্রিম", "Net Payable":"নেট প্রদেয়", "Payroll Month":"বেতন মাস", "Advance Date":"অগ্রিমের তারিখ",
    "Amount (BDT)":"পরিমাণ (টাকা)", "Note":"নোট", "Save":"সংরক্ষণ", "Update":"আপডেট", "Delete":"মুছুন",
    "Search":"অনুসন্ধান", "Refresh":"রিফ্রেশ", "Generate":"তৈরি করুন", "Export Excel":"এক্সেল রপ্তানি", "Export PDF":"PDF রপ্তানি"
}


def is_postgres():
    return bool(DATABASE_URL and not DATABASE_URL.startswith("sqlite://"))


_PG_POOL = None

def _get_pg_pool():
    global _PG_POOL
    if _PG_POOL is None:
        from psycopg2.pool import ThreadedConnectionPool
        minconn = int(os.environ.get("PG_POOL_MIN", "1"))
        maxconn = int(os.environ.get("PG_POOL_MAX", "4"))
        _PG_POOL = ThreadedConnectionPool(
            minconn, maxconn, DATABASE_URL,
            sslmode="require", connect_timeout=10,
            keepalives=1, keepalives_idle=30,
            keepalives_interval=10, keepalives_count=3
        )
    return _PG_POOL


def db_connect():
    if is_postgres():
        return _get_pg_pool().getconn()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def db_release(conn):
    if conn is None:
        return
    if is_postgres():
        try:
            # Never return a connection to the pool with an open transaction.
            try:
                if getattr(conn, "status", None) != 1:  # STATUS_READY = 1
                    conn.rollback()
            except Exception:
                pass
            _get_pg_pool().putconn(conn)
        except Exception:
            try: conn.close()
            except Exception: pass
    else:
        try: conn.close()
        except Exception: pass


def placeholders(sql):
    if is_postgres():
        return sql.replace("?", "%s")
    return sql


def execute(sql, params=(), fetch=False, many=False, commit=False):
    conn = db_connect(); cur = conn.cursor()
    try:
        sql2 = placeholders(sql)
        if many: cur.executemany(sql2, params)
        else: cur.execute(sql2, params)
        rows = cur.fetchall() if fetch else None
        if commit: conn.commit()
        return rows
    finally:
        try: cur.close()
        except Exception: pass
        db_release(conn)


def row_dict(cur, row):
    if row is None: return None
    if hasattr(row, "keys"): return dict(row)
    return {d[0]: row[i] for i, d in enumerate(cur.description)}


def fetch_all(sql, params=()):
    conn = db_connect(); cur = conn.cursor()
    try:
        cur.execute(placeholders(sql), params)
        rows = cur.fetchall()
        return [row_dict(cur, r) for r in rows]
    finally:
        db_release(conn)


def fetch_one(sql, params=()):
    rows = fetch_all(sql, params)
    return rows[0] if rows else None


def scalar(sql, params=(), default=0):
    r = fetch_one(sql, params)
    if not r: return default
    return next(iter(r.values()))


def table_exists(name):
    if is_postgres():
        return bool(scalar("SELECT EXISTS (SELECT 1 FROM information_schema.tables WHERE table_schema='public' AND table_name=?)", (name,), False))
    return bool(scalar("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?", (name,), 0))


def columns(name):
    if not table_exists(name): return set()
    if is_postgres():
        return {r["column_name"] for r in fetch_all("SELECT column_name FROM information_schema.columns WHERE table_schema='public' AND table_name=?", (name,))}
    conn=db_connect(); cur=conn.cursor()
    try:
        cur.execute("PRAGMA table_info(" + name + ")")
        return {r[1] for r in cur.fetchall()}
    finally: db_release(conn)


def add_column(name, col, typ):
    if col in columns(name): return
    execute(f"ALTER TABLE {name} ADD COLUMN {col} {typ}", commit=True)


def hash_password(p):
    return hashlib.sha256(str(p).encode("utf-8")).hexdigest()


def nowstr(): return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def init_db():
    """Create only missing structures. Never drop, truncate, or re-import legacy data."""
    conn=db_connect(); cur=conn.cursor()
    try:
        # Legacy-compatible core schemas.
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS workers (
            id INTEGER PRIMARY KEY, name TEXT NOT NULL, basic_salary REAL NOT NULL DEFAULT 0,
            ot_rate REAL NOT NULL DEFAULT 0, department TEXT, designation TEXT,
            refreshment_bill REAL NOT NULL DEFAULT 0, bangla_name TEXT)"""))
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS attendance (
            id INTEGER PRIMARY KEY, worker_id INTEGER, month_year TEXT NOT NULL,
            present_days INTEGER DEFAULT 0, absent_days INTEGER DEFAULT 0,
            ot_hours REAL DEFAULT 0, advance_deduction REAL DEFAULT 0)"""))
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS daily_attendance (
            id INTEGER PRIMARY KEY, worker_id INTEGER, month_year TEXT NOT NULL,
            day INTEGER NOT NULL, status TEXT NOT NULL DEFAULT 'P')"""))
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS company_settings (
            key TEXT PRIMARY KEY, value TEXT)"""))
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT,
            full_name TEXT, role TEXT DEFAULT 'Operator', active INTEGER DEFAULT 1,
            created_at TEXT, last_login TEXT, password TEXT)"""))
        cur.execute(placeholders("""CREATE TABLE IF NOT EXISTS activity_log (
            id INTEGER PRIMARY KEY, username TEXT, action TEXT, log_time TEXT NOT NULL)"""))
        conn.commit()
    finally: db_release(conn)

    # Existing migrated tables may lack optional web fields. These are additive only.
    for col, typ in [("phone","TEXT"),("address","TEXT"),("joining_date","TEXT"),("status","TEXT")]:
        try: add_column("workers", col, typ)
        except Exception: pass
    try: add_column("users", "password_hash", "TEXT")
    except Exception: pass
    try: add_column("users", "password", "TEXT")
    except Exception: pass

    # Do not overwrite an existing administrator or any migrated record.
    try:
        if not fetch_one("SELECT id FROM users WHERE lower(username)=? LIMIT 1", ("admin",)):
            cols=columns("users")
            vals=[]; names=[]
            for n,v in [("username","admin"),("password_hash",hash_password("admin123")),("password",hash_password("admin123")),("full_name","System Administrator"),("role","Administrator"),("active",1),("created_at",nowstr())]:
                if n in cols: names.append(n); vals.append(v)
            execute("INSERT INTO users("+",".join(names)+") VALUES("+",".join(["?"]*len(vals))+")", vals, commit=True)
    except Exception as e:
        app.logger.warning("admin initialization: %r", e)

    defaults={"company_name":DEFAULT_COMPANY_NAME,"company_address":"","company_phone":"","company_email":"","company_logo":""}
    try:
        for k,v in defaults.items():
            if is_postgres():
                execute("INSERT INTO company_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO NOTHING", (k,v), commit=True)
            else:
                execute("INSERT OR IGNORE INTO company_settings(key,value) VALUES(?,?)", (k,v), commit=True)
    except Exception as e: app.logger.warning("settings initialization: %r", e)


def get_settings():
    try: return {r["key"]: r["value"] for r in fetch_all("SELECT key,value FROM company_settings")}
    except Exception: return {"company_name":DEFAULT_COMPANY_NAME}


def log_activity(username, action):
    try:
        # Actual migrated schema uses log_time, not created_at.
        execute("INSERT INTO activity_log(username,action,log_time) VALUES(?,?,?)", (username,action,nowstr()), commit=True)
    except Exception as e:
        app.logger.warning("Activity log error: %r", e)


def current_user():
    uid=session.get("user_id")
    if not uid: return None
    try: return fetch_one("SELECT * FROM users WHERE id=?", (uid,))
    except Exception: return None


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.path))
        return f(*args, **kwargs)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        u=current_user()
        if not u or str(u.get("role","" )).lower() not in ("administrator","admin"):
            flash("Administrator access required.", "danger"); return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return wrapper


def month_name_year(month=None, year=None):
    month=month or request.values.get("month") or MONTHS[datetime.date.today().month-1]
    year=year or request.values.get("year") or str(datetime.date.today().year)
    if str(month) not in MONTHS:
        month=MONTHS[datetime.date.today().month-1]
    return f"{month} {int(year)}"


def parse_num(v, default=0.0):
    try: return float(v or default)
    except Exception: return default


def legacy_advance_source():
    # Prefer the web table if present; otherwise use the legacy migration table.
    if table_exists("worker_advances"):
        c=columns("worker_advances")
        if {"worker_id","amount"}.issubset(c): return "worker_advances"
    if table_exists("advance_salary"):
        return "advance_salary"
    if table_exists("advances"):
        return "advances"
    return None


def advance_rows(month=None, worker_id=None):
    t=legacy_advance_source()
    if not t: return []
    c=columns(t)
    if t=="worker_advances":
        where=[]; p=[]
        if month: where.append("month_year=?"); p.append(month)
        if worker_id: where.append("worker_id=?"); p.append(worker_id)
        q="SELECT id,worker_id,month_year,advance_date,amount,note FROM worker_advances"
        if where:q+=" WHERE "+" AND ".join(where)
        q+=" ORDER BY id DESC"
        return fetch_all(q,p)
    if t=="advance_salary":
        datecol="date" if "date" in c else ("advance_date" if "advance_date" in c else None)
        q=f"SELECT id,worker_id,{datecol or 'NULL'} AS advance_date,amount,note FROM {t}"
        rows=fetch_all(q)
        for r in rows:r["month_year"]=month_name_from_date(r.get("advance_date"))
        if month: rows=[r for r in rows if r.get("month_year")==month]
        if worker_id: rows=[r for r in rows if str(r.get("worker_id"))==str(worker_id)]
        return rows
    # advances legacy table
    q=f"SELECT id,worker_id,date AS advance_date,amount,note FROM {t}"
    rows=fetch_all(q)
    for r in rows:r["month_year"]=month_name_from_date(r.get("advance_date"))
    if month: rows=[r for r in rows if r.get("month_year")==month]
    if worker_id: rows=[r for r in rows if str(r.get("worker_id"))==str(worker_id)]
    return rows


def month_name_from_date(s):
    try:
        d=str(s)[:10]
        y,m,_=d.split("-")
        return f"{MONTHS[int(m)-1]} {int(y)}"
    except Exception:return ""


def save_advance_record(worker_id, month, adv_date, amount, note):
    t=legacy_advance_source()
    if t=="worker_advances":
        execute("INSERT INTO worker_advances(worker_id,month_year,advance_date,amount,note) VALUES(?,?,?,?,?)", (worker_id,month,adv_date,amount,note), commit=True)
    elif t=="advance_salary":
        c=columns(t); datecol="date" if "date" in c else "advance_date"
        execute(f"INSERT INTO {t}(worker_id,{datecol},amount,note) VALUES(?,?,?,?)", (worker_id,adv_date,amount,note), commit=True)
    elif t=="advances":
        execute("INSERT INTO advances(worker_id,date,amount,note) VALUES(?,?,?,?)", (str(worker_id),adv_date,amount,note), commit=True)
    else:
        # Create only if none exists; this does not touch migrated tables.
        if is_postgres():
            execute("CREATE TABLE IF NOT EXISTS worker_advances (id SERIAL PRIMARY KEY, worker_id INTEGER, month_year TEXT NOT NULL, advance_date TEXT NOT NULL, amount REAL NOT NULL, note TEXT)", commit=True)
        else:
            execute("CREATE TABLE IF NOT EXISTS worker_advances (id INTEGER PRIMARY KEY AUTOINCREMENT, worker_id INTEGER, month_year TEXT NOT NULL, advance_date TEXT NOT NULL, amount REAL NOT NULL, note TEXT)", commit=True)
        execute("INSERT INTO worker_advances(worker_id,month_year,advance_date,amount,note) VALUES(?,?,?,?,?)", (worker_id,month,adv_date,amount,note), commit=True)


def delete_advance_record(aid):
    t=legacy_advance_source()
    if t: execute(f"DELETE FROM {t} WHERE id=?", (aid,), commit=True)


def daily_map(worker_id, month):
    rows=fetch_all("SELECT day,status FROM daily_attendance WHERE worker_id=? AND month_year=?", (worker_id,month))
    return {int(r["day"]): str(r["status"] or "P") for r in rows}


def calculate_salary_bulk(workers, month):
    """Calculate salaries for many workers using a small number of DB queries."""
    ids=[w["id"] for w in workers]
    if not ids: return {}
    att_rows=fetch_all("SELECT * FROM attendance WHERE month_year=? ORDER BY id DESC", (month,))
    att_by={}
    for r in att_rows:
        wid=r.get("worker_id")
        if wid not in att_by: att_by[wid]=r

    daily_rows=fetch_all("SELECT worker_id,day,status FROM daily_attendance WHERE month_year=?", (month,))
    daily_by={}
    for r in daily_rows:
        try: daily_by.setdefault(r.get("worker_id"), {})[int(r.get("day"))]=str(r.get("status") or "P")
        except Exception: pass

    advances_by={}
    t=legacy_advance_source()
    if t:
        c=columns(t)
        try:
            if t=="worker_advances" and {"worker_id","amount"}.issubset(c):
                adv_rows=fetch_all("SELECT worker_id,amount,month_year FROM worker_advances WHERE month_year=?", (month,))
                for r in adv_rows: advances_by[r.get("worker_id")]=advances_by.get(r.get("worker_id"),0)+parse_num(r.get("amount"))
            elif t in ("advance_salary","advances") and "worker_id" in c and "amount" in c:
                datecol="date" if "date" in c else ("advance_date" if "advance_date" in c else None)
                if datecol:
                    adv_rows=fetch_all(f"SELECT worker_id,{datecol} AS advance_date,amount FROM {t}")
                    for r in adv_rows:
                        if month_name_from_date(r.get("advance_date"))==month:
                            wid=r.get("worker_id"); advances_by[wid]=advances_by.get(wid,0)+parse_num(r.get("amount"))
        except Exception as e:
            app.logger.warning("Bulk advance load error: %r", e)

    out={}
    for w in workers:
        att=att_by.get(w["id"], {})
        present=int(att.get("present_days") or 0); absent=int(att.get("absent_days") or 0); ot=parse_num(att.get("ot_hours"),0)
        try:
            mon,ys=month.split(); y=int(ys); mi=MONTHS.index(mon)+1; days=calendar.monthrange(y,mi)[1]
        except Exception:
            days=30; y=datetime.date.today().year; mi=datetime.date.today().month
        basic=parse_num(w.get("basic_salary")); ot_rate=parse_num(w.get("ot_rate")); nasta_rate=parse_num(w.get("refreshment_bill"))
        absent_cut=(basic/days)*absent if days else 0
        earned_basic=max(0,basic-absent_cut)
        ot_amt=ot*ot_rate
        dm=daily_by.get(w["id"],{})
        if dm:
            billable=sum(1 for d in range(1,days+1) if dm.get(d,"P")=="P" and datetime.date(y,mi,d).weekday()!=4)
            nasta=billable*nasta_rate
        else:
            non_friday=sum(1 for d in range(1,days+1) if datetime.date(y,mi,d).weekday()!=4) if days else 0
            nasta=round(present*(non_friday/days)*nasta_rate,2) if days else 0
        advances=advances_by.get(w["id"],0)
        gross=earned_basic+ot_amt+nasta
        out[w["id"]]={"present":present,"absent":absent,"ot":ot,"absent_cut":absent_cut,"earned_basic":earned_basic,"ot_amt":ot_amt,"nasta":nasta,"gross":gross,"advance":advances,"net":gross-advances}
    return out

def calculate_salary(worker, month, att=None):
    if att is None: att=fetch_one("SELECT * FROM attendance WHERE worker_id=? AND month_year=? ORDER BY id DESC LIMIT 1", (worker["id"],month)) or {}
    present=int(att.get("present_days") or 0); absent=int(att.get("absent_days") or 0); ot=parse_num(att.get("ot_hours"),0)
    try:
        mon, ys=month.split(); y=int(ys); mi=MONTHS.index(mon)+1; days=calendar.monthrange(y,mi)[1]
    except Exception: days=30
    # Preserve the legacy payroll logic: absent cut from monthly basic, OT by rate, nasta excluding Fridays.
    basic=parse_num(worker.get("basic_salary")); ot_rate=parse_num(worker.get("ot_rate")); nasta_rate=parse_num(worker.get("refreshment_bill"))
    absent_cut=(basic/days)*absent if days else 0
    earned_basic=max(0,basic-absent_cut)
    ot_amt=ot*ot_rate
    dm=daily_map(worker["id"],month)
    if dm:
        billable=sum(1 for d in range(1,days+1) if dm.get(d,"P")=="P" and datetime.date(y,mi,d).weekday()!=4)
        nasta=billable*nasta_rate
    else:
        non_friday=sum(1 for d in range(1,days+1) if datetime.date(y,mi,d).weekday()!=4) if days else 0
        nasta=round(present*(non_friday/days)*nasta_rate,2) if days else 0
    advances=sum(parse_num(r.get("amount")) for r in advance_rows(month,worker["id"]))
    gross=earned_basic+ot_amt+nasta
    return {"present":present,"absent":absent,"ot":ot,"absent_cut":absent_cut,"earned_basic":earned_basic,"ot_amt":ot_amt,"nasta":nasta,"gross":gross,"advance":advances,"net":gross-advances}


@app.context_processor
def inject_globals():
    settings=get_settings()
    lang=session.get("language","en")
    def tr(s): return LANG.get(s,s) if lang=="bn" else s
    def dept(s): return DEPT_BN.get(str(s),str(s)) if lang=="bn" else str(s or "")
    return {"current_user":current_user(),"settings":settings,"language":lang,"tr":tr,"dept":dept,"display_dept":dept,"months":MONTHS,"years":list(range(2024,2032))}


@app.route("/health")
def health():
    try:
        w=scalar("SELECT COUNT(*) FROM workers",default=0)
        a=scalar("SELECT COUNT(*) FROM attendance",default=0)
        d=scalar("SELECT COUNT(*) FROM daily_attendance",default=0)
        return jsonify({"status":"ok","database":"postgresql" if is_postgres() else "sqlite","workers":w,"attendance":a,"daily_attendance":d})
    except Exception as e:return jsonify({"status":"error","error":repr(e)}),500


@app.route("/db-check")
@login_required
def db_check():
    info={}
    for t in ["workers","attendance","daily_attendance","company_settings","users","activity_log","worker_advances","advance_salary","advances"]:
        try: info[t]={"exists":table_exists(t),"columns":sorted(columns(t)) if table_exists(t) else []}
        except Exception as e: info[t]={"error":repr(e)}
    return render_template("db_check.html",info=info)


@app.route("/login", methods=["GET","POST"])
def login():
    if request.method=="POST":
        u=request.form.get("username","").strip(); p=request.form.get("password","")
        user=fetch_one("SELECT * FROM users WHERE username=? LIMIT 1",(u,))
        ok=False
        if user and int(user.get("active") if user.get("active") is not None else 1):
            ph=user.get("password_hash") or user.get("password") or ""
            ok = ph in (hash_password(p), p)  # supports old migrated password only if it is already a hash/plain legacy value
        if ok:
            session["user_id"]=user["id"]; session["language"]=session.get("language","en")
            try: execute("UPDATE users SET last_login=? WHERE id=?",(nowstr(),user["id"]),commit=True)
            except Exception: pass
            log_activity(u,"Logged in")
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Invalid User ID or Password.","danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    u=current_user();
    if u: log_activity(u.get("username"),"Logged out")
    session.clear(); return redirect(url_for("login"))

@app.route("/language/<lang>")
def language(lang):
    session["language"]="bn" if lang=="bn" else "en"; return redirect(request.referrer or url_for("dashboard"))


@app.route("/")
@login_required
def index(): return redirect(url_for("dashboard"))

@app.route("/dashboard")
@login_required
def dashboard():
    month=month_name_year(); workers=fetch_all("SELECT * FROM workers ORDER BY id")
    salaries=calculate_salary_bulk(workers,month)
    gross=sum(s["gross"] for s in salaries.values())
    adv=sum(s["advance"] for s in salaries.values())
    today=datetime.date.today(); tm=f"{MONTHS[today.month-1]} {today.year}"
    today_rows=fetch_all("SELECT status,COUNT(*) AS c FROM daily_attendance WHERE month_year=? AND day=? GROUP BY status",(tm,today.day))
    mp={r["status"]:r["c"] for r in today_rows}
    depts=fetch_all("SELECT department,COUNT(*) AS worker_count,COALESCE(SUM(basic_salary),0) AS total_basic FROM workers GROUP BY department ORDER BY department")
    return render_template("dashboard.html",workers=workers,month=month,month_name=month.split()[0],year=month.split()[1],gross=gross,advance=adv,today_present=mp.get("P",0),today_absent=mp.get("A",0),dept_rows=depts,total_workers=len(workers))


@app.route("/workers")
@login_required
def workers():
    q=request.args.get("q","").strip().lower(); params=[]
    sql="SELECT * FROM workers"
    if q:
        l=f"%{q}%"; sql+=" WHERE CAST(id AS TEXT) LIKE ? OR lower(name) LIKE ? OR lower(COALESCE(bangla_name,'')) LIKE ? OR lower(COALESCE(department,'')) LIKE ? OR lower(COALESCE(designation,'')) LIKE ?"; params=[l,l,l,l,l]
    sql+=" ORDER BY id DESC"
    return render_template("workers.html",workers=fetch_all(sql,params),q=q)


@app.route("/workers/add", methods=["GET","POST"])
@login_required
def add_worker():
    if request.method=="POST":
        f=request.form; name=f.get("name","").strip()
        if not name: flash("Worker name is required.","danger"); return render_template("worker_form.html",worker=f)
        vals=(name,parse_num(f.get("basic_salary")),parse_num(f.get("ot_rate")),f.get("department","").strip(),f.get("designation","").strip(),parse_num(f.get("refreshment_bill")),f.get("bangla_name","").strip())
        execute("INSERT INTO workers(name,basic_salary,ot_rate,department,designation,refreshment_bill,bangla_name) VALUES(?,?,?,?,?,?,?)",vals,commit=True)
        flash("Worker saved.","success"); return redirect(url_for("workers"))
    return render_template("worker_form.html",worker=None)


@app.route("/workers/edit/<int:worker_id>", methods=["GET","POST"])
@login_required
def edit_worker(worker_id):
    w=fetch_one("SELECT * FROM workers WHERE id=?",(worker_id,))
    if not w: abort(404)
    if request.method=="POST":
        f=request.form
        execute("UPDATE workers SET name=?,basic_salary=?,ot_rate=?,department=?,designation=?,refreshment_bill=?,bangla_name=? WHERE id=?",(f.get("name",""),parse_num(f.get("basic_salary")),parse_num(f.get("ot_rate")),f.get("department",""),f.get("designation",""),parse_num(f.get("refreshment_bill")),f.get("bangla_name",""),worker_id),commit=True)
        flash("Worker updated.","success"); return redirect(url_for("workers"))
    return render_template("worker_form.html",worker=w)


@app.route("/workers/delete/<int:worker_id>", methods=["POST","GET"])
@login_required
def delete_worker(worker_id):
    execute("DELETE FROM workers WHERE id=?",(worker_id,),commit=True)
    # Explicit user action only; never used during initialization/migration.
    for t in ["attendance","daily_attendance","worker_advances","advance_salary","advances"]:
        if table_exists(t):
            try: execute(f"DELETE FROM {t} WHERE worker_id=?",(worker_id,),commit=True)
            except Exception: pass
    flash("Worker deleted.","success"); return redirect(url_for("workers"))


@app.route("/attendance")
@login_required
def attendance():
    req_month=request.args.get("month")
    req_year=request.args.get("year")
    month=month_name_year(req_month, req_year) if req_month and req_year else month_name_year()
    wid=request.args.get("worker_id",""); w=None; days=[]; summary={}
    if wid:
        w=fetch_one("SELECT * FROM workers WHERE id=?",(wid,))
        if w:
            dm=daily_map(w["id"],month); mon,ys=month.split(); y=int(ys); mi=MONTHS.index(mon)+1; nd=calendar.monthrange(y,mi)[1]
            days=[{"day":d,"status":dm.get(d,"P"),"date":datetime.date(y,mi,d)} for d in range(1,nd+1)]
            summary=calculate_salary(w,month)
    return render_template("attendance.html",workers=fetch_all("SELECT id,name,department FROM workers ORDER BY id"),worker=w,days=days,summary=summary,month=month)


@app.route("/attendance/save", methods=["POST"])
@login_required
def save_attendance():
    f=request.form; wid=int(f.get("worker_id")); month=month_name_year(f.get("month"),f.get("year")) if f.get("year") else f.get("month_year") or month_name_year()
    present=int(f.get("present_days") or 0); absent=int(f.get("absent_days") or 0); ot=parse_num(f.get("ot_hours"))
    dm={}
    raw=f.get("statuses_json")
    if raw:
        try: import json; dm=json.loads(raw)
        except Exception: dm={}
    # Also accept fields day_1=P etc.
    for k,v in f.items():
        if k.startswith("day_"): dm[k[4:]]=v
    for d,st in dm.items():
        try:
            day=int(d); st="A" if str(st).upper().startswith("A") else "P"
            existing_day = fetch_one("SELECT id FROM daily_attendance WHERE worker_id=? AND month_year=? AND day=? ORDER BY id DESC LIMIT 1", (wid, month, day))
            if existing_day:
                execute("UPDATE daily_attendance SET status=? WHERE id=?", (st, existing_day["id"]), commit=True)
            else:
                execute("INSERT INTO daily_attendance(worker_id,month_year,day,status) VALUES(?,?,?,?)", (wid,month,day,st), commit=True)
        except Exception: pass
    existing=fetch_one("SELECT id FROM attendance WHERE worker_id=? AND month_year=? ORDER BY id DESC LIMIT 1",(wid,month))
    if existing:
        execute("UPDATE attendance SET present_days=?,absent_days=?,ot_hours=? WHERE id=?",(present,absent,ot,existing["id"]),commit=True)
    else:
        acols = columns("attendance")
        names = ["worker_id","month_year","present_days","absent_days","ot_hours"]
        vals = [wid,month,present,absent,ot]
        if "advance_deduction" in acols:
            names.append("advance_deduction"); vals.append(0)
        execute("INSERT INTO attendance(" + ",".join(names) + ") VALUES(" + ",".join(["?"]*len(vals)) + ")", vals, commit=True)
    flash("Attendance saved.","success"); return redirect(url_for("attendance",worker_id=wid,month=month.split()[0],year=month.split()[1]))


@app.route("/payslip")
@login_required
def payslip():
    month=month_name_year(); wid=request.args.get("worker_id"); w=fetch_one("SELECT * FROM workers WHERE id=?",(wid,)) if wid else None
    salary=calculate_salary(w,month) if w else {}
    return render_template("payslip.html",workers=fetch_all("SELECT id,name,department FROM workers ORDER BY id"),worker=w,summary=salary,month=month)


@app.route("/department")
@login_required
def department():
    month=month_name_year(); dept=request.args.get("department","")
    ws=fetch_all("SELECT * FROM workers" + (" WHERE department=?" if dept else "") + " ORDER BY id",(dept,) if dept else ())
    salary_map=calculate_salary_bulk(ws,month)
    rows=[{**w,**salary_map.get(w["id"],{})} for w in ws]
    return render_template("department.html",rows=rows,month=month,department=dept,departments=[r["department"] for r in fetch_all("SELECT DISTINCT department FROM workers WHERE department IS NOT NULL AND department<>'' ORDER BY department")])

# Compatibility endpoint used by the current dashboard template.
# The existing Department Salary page is the report view.
app.add_url_rule("/report", endpoint="report", view_func=department, methods=["GET"])


@app.route("/advance")
@app.route("/advances")
@login_required
def advance():
    month=month_name_year(); wid=request.args.get("worker_id")
    return render_template("advance.html",workers=fetch_all("SELECT id,name,department FROM workers ORDER BY id"),rows=advance_rows(month,wid),month=month)

# Compatibility endpoints for older templates.
# Flask endpoint names come from function names, so the route aliases above
# do not create an endpoint named "advances".
app.add_url_rule("/advances", endpoint="advances", view_func=advance, methods=["GET"])



@app.route("/advance/save", methods=["POST"])
@app.route("/advances/save", methods=["POST"])
@login_required
def save_advance():
    f=request.form; wid=int(f.get("worker_id")); month=f.get("month_year") or month_name_year(f.get("month"),f.get("year")); adv_date=f.get("advance_date") or datetime.date.today().isoformat(); amount=parse_num(f.get("amount")); note=f.get("note","")
    if amount<=0: flash("Amount must be greater than zero.","danger"); return redirect(url_for("advance",month=month))
    save_advance_record(wid,month,adv_date,amount,note); flash("Advance saved.","success"); return redirect(url_for("advance",month=month))


@app.route("/advance/delete/<int:aid>")
@app.route("/advances/delete/<int:aid>")
@login_required
def delete_advance(aid):
    delete_advance_record(aid); flash("Advance deleted.","success"); return redirect(url_for("advance"))


@app.route("/settings", methods=["GET","POST"])
@admin_required
def settings():
    if request.method=="POST":
        for k in ["company_name","company_address","company_phone","company_email","company_logo"]:
            v=request.form.get(k,"")
            if is_postgres(): execute("INSERT INTO company_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=EXCLUDED.value",(k,v),commit=True)
            else: execute("INSERT OR REPLACE INTO company_settings(key,value) VALUES(?,?)",(k,v),commit=True)
        flash("Settings saved.","success")
    return render_template("settings.html",settings=get_settings())


@app.route("/users")
@admin_required
def users(): return render_template("users.html",users=fetch_all("SELECT id,username,full_name,role,active,created_at,last_login FROM users ORDER BY id"))

@app.route("/users/add", methods=["GET","POST"])
@admin_required
def add_user():
    if request.method=="POST":
        f=request.form; username=f.get("username","").strip(); pw=f.get("password","")
        if not username or not pw: flash("Username and password are required.","danger"); return render_template("user_form.html",user=f)
        ph=hash_password(pw); cols=columns("users")
        names=[]; vals=[]
        for n,v in [("username",username),("password_hash",ph),("password",ph),("full_name",f.get("full_name","")),("role",f.get("role","Operator")),("active",1),("created_at",nowstr())]:
            if n in cols:names.append(n);vals.append(v)
        try: execute("INSERT INTO users("+",".join(names)+") VALUES("+",".join(["?"]*len(vals))+")",vals,commit=True); flash("User created.","success"); return redirect(url_for("users"))
        except Exception as e: flash(f"Could not create user: {e}","danger")
    return render_template("user_form.html",user=None)

@app.route("/users/delete/<int:user_id>")
@admin_required
def delete_user(user_id):
    if current_user() and current_user().get("id")==user_id: flash("You cannot delete the logged-in user.","danger")
    else: execute("DELETE FROM users WHERE id=?",(user_id,),commit=True); flash("User deleted.","success")
    return redirect(url_for("users"))

@app.route("/activity")
@admin_required
def activity(): return render_template("activity.html",logs=fetch_all("SELECT * FROM activity_log ORDER BY id DESC LIMIT 500"))


def export_rows(rows, filename, title="Reedoy Payroll"):
    if openpyxl is None: abort(503,description="openpyxl is not installed")
    wb=Workbook(); ws=wb.active; ws.title="Payroll"
    if rows:
        headers=list(rows[0].keys()); ws.append(headers)
        for r in rows: ws.append([r.get(h) for h in headers])
    else: ws.append(["No records"])
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=filename,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/export/workers.xlsx")
@login_required
def export_workers(): return export_rows(fetch_all("SELECT * FROM workers ORDER BY id"),"reedoy_workers.xlsx")

@app.route("/export/advances.xlsx")
@login_required
def export_advances(): return export_rows(advance_rows(month_name_year()),"reedoy_advances.xlsx")

@app.route("/export/department.xlsx")
@login_required
def export_department():
    month=month_name_year(); dept=request.args.get("department",""); ws=fetch_all("SELECT * FROM workers"+(" WHERE department=?" if dept else "")+" ORDER BY id",(dept,) if dept else ())
    salary_map=calculate_salary_bulk(ws,month)
    rows=[{"ID":w["id"],"Worker Name":w["name"],"Department":w.get("department"),"Designation":w.get("designation"),**salary_map.get(w["id"],{})} for w in ws]
    return export_rows(rows,"reedoy_department_salary.xlsx")

@app.route("/export/payslip.pdf")
@login_required
def export_payslip_pdf():
    if SimpleDocTemplate is None: abort(503,description="reportlab is not installed")
    wid=request.args.get("worker_id"); month=month_name_year(); w=fetch_one("SELECT * FROM workers WHERE id=?",(wid,))
    if not w: abort(404)
    s=calculate_salary(w,month); bio=io.BytesIO(); doc=SimpleDocTemplate(bio,pagesize=A4,rightMargin=35,leftMargin=35,topMargin=35,bottomMargin=35)
    styles=getSampleStyleSheet(); story=[Paragraph(DEFAULT_COMPANY_NAME,styles["Title"]),Paragraph("Payroll Payslip - "+month,styles["Heading2"]),Spacer(1,12)]
    data=[["Worker ID",w["id"]],["Worker Name",w["name"]],["Department",w.get("department") or ""],["Basic Salary",f"BDT {s['earned_basic']:,.2f}"],["Absent Deduction",f"BDT {s['absent_cut']:,.2f}"],["OT",f"BDT {s['ot_amt']:,.2f}"],["Nasta",f"BDT {s['nasta']:,.2f}"],["Gross",f"BDT {s['gross']:,.2f}"],["Advance",f"BDT {s['advance']:,.2f}"],["Net Payable",f"BDT {s['net']:,.2f}"]]
    story.append(Table(data,colWidths=[180,280],style=TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.grey),("FONTNAME",(0,0),(-1,-1),"Helvetica"),("PADDING",(0,0),(-1,-1),6)])))
    doc.build(story); bio.seek(0); return send_file(bio,as_attachment=True,download_name=f"payslip_{w['id']}_{month.replace(' ','_')}.pdf",mimetype="application/pdf")

@app.route("/export/workers.csv")
@login_required
def export_workers_csv():
    rows=fetch_all("SELECT * FROM workers ORDER BY id"); bio=io.StringIO();
    if rows:
        wr=csv.DictWriter(bio,fieldnames=list(rows[0].keys()));wr.writeheader();wr.writerows(rows)
    return send_file(io.BytesIO(bio.getvalue().encode("utf-8-sig")),as_attachment=True,download_name="reedoy_workers.csv",mimetype="text/csv")


@app.errorhandler(404)
def not_found(e):
    try:return render_template("404.html"),404
    except Exception:return "404 - Page not found",404

@app.errorhandler(500)
def server_error(e):
    app.logger.exception("Unhandled application error")
    try:return render_template("500.html",error=e),500
    except Exception:return f"500 - Internal Server Error\n{e}",500


# Initialize on import so gunicorn app:app can start.
try:
    init_db()
except Exception as e:
    app.logger.exception("Database initialization error: %r", e)

if __name__ == "__main__":
    app.run(host="0.0.0.0",port=int(os.environ.get("PORT",5000)),debug=False)
